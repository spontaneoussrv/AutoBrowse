"""
AutoBrowse
Developed by Sourab Kumar Saha

Sits in the system tray, starts with Windows, opens your chosen websites in
Chrome, refreshes them on a fixed or random schedule and browses inner pages
at a natural, human pace (curved mouse paths, wheel scrolling, reading pauses).
"""

import json
import logging
import math
import os
import queue
import random
import re
import difflib
import socket
import subprocess
import sys
import threading
import time
import urllib.parse
import urllib.request
import uuid
import webbrowser
from collections import deque
from datetime import datetime
from logging.handlers import RotatingFileHandler

import tkinter as tk
from tkinter import ttk, messagebox, filedialog

REQUIRED = {"pystray": "pystray", "websockets": "websockets", "PIL": "Pillow", "customtkinter": "customtkinter",
            "openpyxl": "openpyxl"}


def ensure_packages():
    """Install any missing library automatically on first run."""
    import importlib
    missing = []
    for mod, pkg in REQUIRED.items():
        try:
            importlib.import_module(mod)
        except ImportError:
            missing.append(pkg)
    if not missing or getattr(sys, "frozen", False):
        return
    splash = tk.Tk()
    splash.title("AutoBrowse")
    splash.geometry("380x110")
    splash.resizable(False, False)
    tk.Label(splash, text="Setting up AutoBrowse for the first time...\nInstalling: " + ", ".join(missing),
             font=("Segoe UI", 10), justify="center").pack(expand=True)
    splash.update()
    py = sys.executable
    if py.lower().endswith("pythonw.exe"):
        py = py[:-len("pythonw.exe")] + "python.exe"
    flags = 0x08000000 if os.name == "nt" else 0
    result = subprocess.run([py, "-m", "pip", "install", "--disable-pip-version-check", "--quiet"] + missing,
                            creationflags=flags, capture_output=True, text=True)
    splash.destroy()
    if result.returncode != 0:
        messagebox.showerror("AutoBrowse", "Could not install: %s\n\n%s" % (", ".join(missing),
                                                                           (result.stderr or "")[-600:]))
        sys.exit(1)
    importlib.invalidate_caches()


ensure_packages()

import customtkinter as ctk
import pystray
from websockets.sync.server import serve as ws_serve
from PIL import Image, ImageDraw, ImageTk


APP_NAME = "AutoBrowse"
APP_DIR = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), APP_NAME)
CONFIG_PATH = os.path.join(APP_DIR, "settings.json")
LOG_PATH = os.path.join(APP_DIR, "activity.log")
HELPER_DIR = os.path.join(APP_DIR, "Chrome Helper")
INSTANCE_PORT = 47831
BRIDGE_PORT = 47832

DEFAULT_SKIP = (
    "logout, log-out, log out, signout, sign-out, sign out, delete, remove, unsubscribe, "
    "cart, checkout, payment, pay now, buy now, download, mailto:, tel:, whatsapp"
)

DEFAULT_TARGETS = ""

DEFAULT_CONFIG = {
    "chrome_path": "",
    "profile": "Default",
    "theme": "System",
    "pause_busy": True,
    "busy_minutes": 2,
    "pause_battery": False,
    "startup": True,
    "hours_on": False,
    "hours_from": "09:00",
    "hours_to": "23:00",
    "sites": [],
}


def new_site(url=""):
    return {
        "id": uuid.uuid4().hex,
        "url": url,
        "enabled": True,
        "mode": "random",
        "fixed_min": 5.0,
        "rand_min": 3.0,
        "rand_max": 8.0,
        "browse": True,
        "pages_min": 1,
        "pages_max": 3,
        "stay_min": 15,
        "stay_max": 45,
        "scroll": True,
        "home": True,
        "front": True,
        "skip": DEFAULT_SKIP,
        "targets": DEFAULT_TARGETS,
        "real_mouse": False,
        "alert_change": True,
        "alert_logout": True,
        "alert_down": True,
        "hours_on": False,
        "hours_from": "09:00",
        "hours_to": "23:00",
        "limit_on": False,
        "limit_visits": 50,
        "limit_minutes": 60,
    }


# ----------------------------------------------------------------------------
# Settings and logging
# ----------------------------------------------------------------------------

def load_config():
    cfg = dict(DEFAULT_CONFIG)
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            cfg.update(json.load(f))
    except (OSError, ValueError):
        pass
    sites = []
    for s in cfg.get("sites", []):
        base = new_site()
        base.update(s)
        if not cfg.get("targets_default_set") and not base.get("targets"):
            base["targets"] = DEFAULT_TARGETS
        sites.append(base)
    cfg["sites"] = sites
    cfg["targets_default_set"] = True
    return cfg


def save_config(cfg):
    os.makedirs(APP_DIR, exist_ok=True)
    tmp = CONFIG_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2)
    os.replace(tmp, CONFIG_PATH)


LOG_LINES = deque(maxlen=600)


class MemoryHandler(logging.Handler):
    def emit(self, record):
        LOG_LINES.append(self.format(record))


def setup_logging():
    os.makedirs(APP_DIR, exist_ok=True)
    log = logging.getLogger(APP_NAME)
    log.setLevel(logging.INFO)
    fmt = logging.Formatter("%(asctime)s   %(message)s", "%d %b %I:%M:%S %p")
    fh = RotatingFileHandler(LOG_PATH, maxBytes=1_000_000, backupCount=2, encoding="utf-8")
    fh.setFormatter(fmt)
    mh = MemoryHandler()
    mh.setFormatter(fmt)
    log.addHandler(fh)
    log.addHandler(mh)
    return log


log = setup_logging()


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------

def normalize_url(u):
    u = (u or "").strip()
    if u and "://" not in u:
        u = "https://" + u
    return u


def host_of(u):
    try:
        h = (urllib.parse.urlparse(u).hostname or "").lower()
    except ValueError:
        return ""
    return h[4:] if h.startswith("www.") else h


def to_num(value, fallback, low=0.0, as_int=False):
    try:
        n = float(str(value).strip())
    except ValueError:
        n = fallback
    n = max(low, n)
    return int(round(n)) if as_int else n


def clock(text, fallback):
    try:
        return datetime.strptime(text.strip(), "%H:%M").strftime("%H:%M")
    except ValueError:
        return fallback


def nice(n):
    return str(int(n)) if float(n).is_integer() else ("%g" % n)


def schedule_text(s):
    if not s.get("enabled"):
        return "Off"
    if s["mode"] == "fixed":
        return "Every %s min" % nice(s["fixed_min"])
    return "%s to %s min" % (nice(s["rand_min"]), nice(s["rand_max"]))


def within_hours(cfg):
    if not cfg.get("hours_on"):
        return True
    try:
        a = datetime.strptime(cfg["hours_from"].strip(), "%H:%M").time()
        b = datetime.strptime(cfg["hours_to"].strip(), "%H:%M").time()
    except (ValueError, KeyError):
        return True
    now = datetime.now().time()
    return a <= now <= b if a <= b else (now >= a or now <= b)


def launch_command():
    if getattr(sys, "frozen", False):
        return '"%s" --tray' % sys.executable
    exe = sys.executable
    w = os.path.join(os.path.dirname(exe), "pythonw.exe")
    if os.path.exists(w):
        exe = w
    return '"%s" "%s" --tray' % (exe, os.path.abspath(__file__))


def set_startup(enabled):
    if os.name != "nt":
        return
    import winreg
    path = r"Software\Microsoft\Windows\CurrentVersion\Run"
    with winreg.OpenKey(winreg.HKEY_CURRENT_USER, path, 0, winreg.KEY_SET_VALUE) as key:
        if enabled:
            winreg.SetValueEx(key, APP_NAME, 0, winreg.REG_SZ, launch_command())
        else:
            try:
                winreg.DeleteValue(key, APP_NAME)
            except FileNotFoundError:
                pass


def make_shortcuts():
    """Desktop and Start Menu shortcuts with the green AutoBrowse icon."""
    if os.name != "nt":
        return
    import shutil
    ico = os.path.join(APP_DIR, "AutoBrowse.ico")
    try:
        if os.path.exists(resource("AutoBrowse.ico")):
            shutil.copyfile(resource("AutoBrowse.ico"), ico)
        else:
            make_icon(256).save(ico, sizes=[(16, 16), (24, 24), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)])
    except OSError:
        return
    if getattr(sys, "frozen", False):
        target, args = sys.executable, ""
    else:
        target = sys.executable
        w = os.path.join(os.path.dirname(target), "pythonw.exe")
        if os.path.exists(w):
            target = w
        args = '"%s"' % os.path.abspath(__file__)
    q = lambda v: v.replace("'", "''")
    script = (
        "$w = New-Object -ComObject WScript.Shell;"
        "foreach ($d in @([Environment]::GetFolderPath('Desktop'), "
        "(Join-Path $env:APPDATA 'Microsoft\\Windows\\Start Menu\\Programs'))) {"
        "$s = $w.CreateShortcut((Join-Path $d 'AutoBrowse.lnk'));"
        "$s.TargetPath = '%s'; $s.Arguments = '%s'; $s.WorkingDirectory = '%s';"
        "$s.IconLocation = '%s,0'; $s.Description = 'AutoBrowse'; $s.Save() }"
    ) % (q(target), q(args), q(os.path.dirname(os.path.abspath(sys.argv[0]))), q(ico))
    try:
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
                       capture_output=True, creationflags=0x08000000, timeout=30)
    except Exception:
        pass


def find_chrome(custom=""):
    if custom and os.path.isfile(custom):
        return custom
    found = []
    for env in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        base = os.environ.get(env)
        if base:
            found.append(os.path.join(base, "Google", "Chrome", "Application", "chrome.exe"))
    if os.name == "nt":
        import winreg
        for hive in (winreg.HKEY_CURRENT_USER, winreg.HKEY_LOCAL_MACHINE):
            try:
                with winreg.OpenKey(hive, r"Software\Microsoft\Windows\CurrentVersion\App Paths\chrome.exe") as k:
                    found.append(winreg.QueryValue(k, None))
            except OSError:
                pass
    found += [
        "/usr/bin/google-chrome",
        "/usr/bin/chromium",
        "/opt/pw-browsers/chromium",
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    ]
    for c in found:
        if c and os.path.isfile(c):
            return c
    return None


# ----------------------------------------------------------------------------
# Chrome control
# ----------------------------------------------------------------------------

class Stopped(Exception):
    pass


# ----------------------------------------------------------------------------
# Chrome helper extension (installed once into your own Chrome profile)
# ----------------------------------------------------------------------------

HELPER_MANIFEST = {
    "manifest_version": 3,
    "name": "AutoBrowse Helper",
    "version": "1.1",
    "description": "Lets the AutoBrowse app refresh and browse the websites you choose.",
    "permissions": ["tabs", "debugger", "alarms"],
    "background": {"service_worker": "background.js"},
    "icons": {"16": "icon16.png", "48": "icon48.png", "128": "icon128.png"},
}

HELPER_JS = r"""
const PORT = BRIDGE_PORT;
let ws = null;
let keep = null;

function connect() {
  if (ws && (ws.readyState === 0 || ws.readyState === 1)) return;
  try { ws = new WebSocket('ws://127.0.0.1:' + PORT); } catch (e) { ws = null; return; }
  ws.onopen = () => {
    ws.send(JSON.stringify({type: 'hello', version: chrome.runtime.getManifest().version}));
    clearInterval(keep);
    keep = setInterval(() => { if (ws && ws.readyState === 1) ws.send('{"type":"ping"}'); }, 20000);
  };
  ws.onmessage = ev => handle(JSON.parse(ev.data));
  ws.onclose = () => { clearInterval(keep); ws = null; setTimeout(connect, 5000); };
  ws.onerror = () => {};
}

async function handle(msg) {
  let reply;
  try { reply = {id: msg.id, ok: true, result: await run(msg)}; }
  catch (e) { reply = {id: msg.id, ok: false, error: String((e && e.message) || e)}; }
  if (ws && ws.readyState === 1) ws.send(JSON.stringify(reply));
}

async function attach(tabId) {
  try { await chrome.debugger.attach({tabId}, '1.3'); }
  catch (e) { if (!String(e.message || e).includes('already attached')) throw e; }
}

async function run(m) {
  switch (m.cmd) {
    case 'tabs': {
      const tabs = await chrome.tabs.query({});
      return tabs.map(t => ({id: t.id, url: t.url || t.pendingUrl || '', windowId: t.windowId, active: t.active}));
    }
    case 'create': {
      const wins = await chrome.windows.getAll({windowTypes: ['normal']});
      if (!wins.length) {
        const w = await chrome.windows.create({url: m.url, focused: !!m.active});
        return {id: w.tabs[0].id, url: m.url};
      }
      const t = await chrome.tabs.create({url: m.url, active: !!m.active});
      return {id: t.id, url: m.url};
    }
    case 'activate': await chrome.tabs.update(m.tabId, {active: true}); return true;
    case 'navigate': await chrome.tabs.update(m.tabId, {url: m.url}); return true;
    case 'focus': {
      const t = await chrome.tabs.update(m.tabId, {active: true});
      const w = await chrome.windows.get(t.windowId);
      if (w.state === 'minimized') await chrome.windows.update(w.id, {state: 'normal'});
      await chrome.windows.update(w.id, {focused: true});
      return true;
    }
    case 'attach': await attach(m.tabId); return true;
    case 'detach': try { await chrome.debugger.detach({tabId: m.tabId}); } catch (e) {} return true;
    case 'cdp':
      try { return await chrome.debugger.sendCommand({tabId: m.tabId}, m.method, m.params || {}); }
      catch (e) {
        if (!String(e.message || e).includes('not attached')) throw e;
        await attach(m.tabId);
        return await chrome.debugger.sendCommand({tabId: m.tabId}, m.method, m.params || {});
      }
  }
  throw new Error('Unknown command');
}

chrome.alarms.create('keep-connected', {periodInMinutes: 0.5});
chrome.alarms.onAlarm.addListener(connect);
chrome.runtime.onStartup.addListener(connect);
chrome.runtime.onInstalled.addListener(connect);
connect();
"""


def write_helper():
    os.makedirs(HELPER_DIR, exist_ok=True)
    with open(os.path.join(HELPER_DIR, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(HELPER_MANIFEST, f, indent=2)
    with open(os.path.join(HELPER_DIR, "background.js"), "w", encoding="utf-8") as f:
        f.write(HELPER_JS.replace("BRIDGE_PORT", str(BRIDGE_PORT)).lstrip())
    for size in (16, 48, 128):
        make_icon(size).save(os.path.join(HELPER_DIR, "icon%d.png" % size))


class Bridge:
    """Local connection between AutoBrowse and the helper extension inside Chrome."""

    def __init__(self):
        self.conn = None
        self.send_lock = threading.Lock()
        self.pending = {}
        self.counter = 0
        self.ready = threading.Event()
        self.version = ""

    def start(self):
        threading.Thread(target=self.serve, daemon=True).start()

    def serve(self):
        try:
            with ws_serve(self.handler, "127.0.0.1", BRIDGE_PORT, ping_interval=None, max_size=2 ** 24) as server:
                server.serve_forever()
        except OSError as e:
            log.info("Could not open the Chrome helper connection: %s", e)

    def handler(self, ws):
        origin = ws.request.headers.get("Origin", "") if ws.request else ""
        if not origin.startswith("chrome-extension://"):
            ws.close()
            return
        self.conn = ws
        try:
            for raw in ws:
                msg = json.loads(raw)
                if msg.get("type") == "hello":
                    self.version = str(msg.get("version", ""))
                    if not self.ready.is_set():
                        log.info("Chrome helper connected")
                    self.ready.set()
                elif "id" in msg:
                    slot = self.pending.get(msg["id"])
                    if slot:
                        slot["reply"] = msg
                        slot["event"].set()
        except Exception:
            pass
        finally:
            if self.conn is ws:
                self.conn = None
                self.ready.clear()
                log.info("Chrome helper disconnected")

    def connected(self):
        return self.conn is not None and self.ready.is_set()

    def call(self, cmd, timeout=30, **kw):
        ws = self.conn
        if ws is None:
            raise RuntimeError("Chrome helper is not connected")
        with self.send_lock:
            self.counter += 1
            mid = self.counter
            slot = {"event": threading.Event(), "reply": None}
            self.pending[mid] = slot
            kw.update(id=mid, cmd=cmd)
            ws.send(json.dumps(kw))
        try:
            if not slot["event"].wait(timeout):
                raise TimeoutError("Chrome did not answer (%s)" % cmd)
            reply = slot["reply"]
            if not reply.get("ok"):
                raise RuntimeError(reply.get("error") or "Chrome error")
            return reply.get("result")
        finally:
            self.pending.pop(mid, None)


BRIDGE = Bridge()


def chrome_profiles():
    """Return [(folder, display name)] for the Chrome profiles on this PC."""
    path = os.path.join(os.environ.get("LOCALAPPDATA", ""), "Google", "Chrome", "User Data", "Local State")
    try:
        with open(path, "r", encoding="utf-8") as f:
            cache = json.load(f)["profile"]["info_cache"]
        items = [(k, (v.get("name") or k)) for k, v in cache.items()]
        items.sort(key=lambda kv: (kv[0] != "Default", kv[0]))
        return items or [("Default", "Default")]
    except (OSError, ValueError, KeyError):
        return [("Default", "Default")]


class Chrome:
    def __init__(self, cfg):
        self.path = cfg.get("chrome_path", "")
        self.profile = cfg.get("profile") or "Default"
        self.bridge = BRIDGE

    def alive(self):
        return self.bridge.connected()

    def launch(self, first_url=None, wait=True):
        exe = find_chrome(self.path)
        if not exe:
            raise RuntimeError("Google Chrome was not found. Set its location in the General tab.")
        args = [exe, "--profile-directory=%s" % self.profile, "--silent-debugger-extension-api"]
        if first_url:
            args.append(first_url)
        flags = 0x00000008 if os.name == "nt" else 0
        subprocess.Popen(args, creationflags=flags, close_fds=True,
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if wait and not self.bridge.ready.wait(45):
            raise RuntimeError("Chrome opened, but the AutoBrowse Helper is not connected. "
                               "Install it from the General tab (one time only).")

    def tabs(self):
        return self.bridge.call("tabs") or []

    def new_tab(self, url, front=False):
        return self.bridge.call("create", url=url, active=front)

    def focus(self, tab_id):
        self.bridge.call("focus", tabId=tab_id)

    def activate(self, tab_id):
        try:
            self.bridge.call("activate", tabId=tab_id)
        except Exception:
            pass


class Tab:
    def __init__(self, target):
        self.id = target["id"]
        self.bridge = BRIDGE
        self.bridge.call("attach", tabId=self.id)
        self.ctx = None
        self.mx = random.uniform(250, 700)
        self.my = random.uniform(180, 420)

    def close(self):
        try:
            self.bridge.call("detach", timeout=5, tabId=self.id)
        except Exception:
            pass

    def send(self, method, params=None, timeout=30):
        return self.bridge.call("cdp", timeout=timeout, tabId=self.id, method=method, params=params or {}) or {}

    def new_world(self):
        """A private script context: shares the page, invisible to the page's own scripts."""
        frame = self.send("Page.getFrameTree")["frameTree"]["frame"]["id"]
        r = self.send("Page.createIsolatedWorld", {"frameId": frame, "worldName": APP_NAME})
        self.ctx = r["executionContextId"]

    def js(self, expression, timeout=20):
        for attempt in range(3):
            if self.ctx is None:
                try:
                    self.new_world()
                except Exception:
                    time.sleep(0.4)
                    continue
            params = {"expression": expression, "returnByValue": True, "awaitPromise": True,
                      "contextId": self.ctx}
            try:
                r = self.send("Runtime.evaluate", params, timeout)
            except RuntimeError as e:
                if "context" in str(e).lower():
                    self.ctx = None
                    continue
                raise
            return r.get("result", {}).get("value")
        raise RuntimeError("Page is not ready")


HOVER_JS = r"""
(() => {
  const vis = [...document.querySelectorAll('a[href], button')].filter(a => {
    const r = a.getBoundingClientRect();
    return r.width > 8 && r.height > 8 && r.top > 40 && r.bottom < innerHeight - 20 && r.left >= 0 && r.right <= innerWidth;
  });
  if (!vis.length) return null;
  const r = vis[Math.floor(Math.random() * vis.length)].getBoundingClientRect();
  return JSON.stringify({x: r.left + r.width * (0.3 + Math.random() * 0.4), y: r.top + r.height * (0.35 + Math.random() * 0.3)});
})()
"""


LINKS_JS = r"""
(() => {
  const skip = SKIP_LIST;
  const norm = h => h.replace(/^www\./, '');
  const host = norm(location.hostname);
  const here = location.href.split('#')[0];
  const fileExt = /\.(pdf|zip|rar|7z|exe|msi|dmg|apk|jpe?g|png|gif|webp|svg|mp4|mp3|wav|xlsx?|xlsm|docx?|pptx?|csv)$/i;
  const seen = new Set(); const els = []; const out = [];
  for (const a of document.querySelectorAll('a[href]')) {
    let u; try { u = new URL(a.getAttribute('href'), location.href); } catch (e) { continue; }
    if (u.protocol !== 'http:' && u.protocol !== 'https:') continue;
    if (norm(u.hostname) !== host) continue;
    if (a.hasAttribute('download') || fileExt.test(u.pathname)) continue;
    const clean = u.href.split('#')[0];
    if (clean === here || seen.has(clean)) continue;
    const text = (a.innerText || a.getAttribute('aria-label') || a.title || '').replace(/\s+/g, ' ').trim();
    const hay = (clean + ' ' + text).toLowerCase();
    if (skip.some(w => w && hay.includes(w))) continue;
    const r = a.getBoundingClientRect();
    if (r.width < 6 || r.height < 6) continue;
    const st = getComputedStyle(a);
    if (st.visibility === 'hidden' || st.display === 'none' || parseFloat(st.opacity) < 0.1) continue;
    seen.add(clean);
    out.push({i: els.length, href: clean, text: text.slice(0, 70)});
    els.push(a);
  }
  window.__abLinks = els;
  return JSON.stringify(out);
})()
"""

RECT_JS = r"""
(() => {
  const a = window.__abLinks && window.__abLinks[INDEX];
  if (!a || !a.isConnected) return null;
  const r = a.getBoundingClientRect();
  const cx = r.left + r.width * FX, cy = r.top + r.height * FY;
  let ok = cx > 0 && cy > 0 && cx < innerWidth && cy < innerHeight;
  if (ok) { const e = document.elementFromPoint(cx, cy); ok = !!e && (e === a || a.contains(e) || e.closest('a') === a); }
  return JSON.stringify({top: r.top, h: r.height, cx: cx, cy: cy, ok: ok, vh: innerHeight});
})()
"""

VIEW_JS = r"""
(() => {
  const se = document.scrollingElement || document.documentElement;
  const hidden = e => e && /hidden|clip/.test(getComputedStyle(e).overflowY);
  const pick = () => {
    if (se.scrollHeight > innerHeight + 40 && !hidden(document.documentElement) && !hidden(document.body)) return se;
    let best = null, area = 0;
    for (const e of document.querySelectorAll('*')) {
      if (e.scrollHeight <= e.clientHeight + 40) continue;
      if (!/auto|scroll|overlay/.test(getComputedStyle(e).overflowY)) continue;
      const r = e.getBoundingClientRect();
      const a = Math.max(0, Math.min(r.right, innerWidth) - Math.max(r.left, 0)) *
                Math.max(0, Math.min(r.bottom, innerHeight) - Math.max(r.top, 0));
      if (a > area) { area = a; best = e; }
    }
    return best || se;
  };
  let el = window.__abScroller;
  if (!el || !el.isConnected || el.scrollHeight <= el.clientHeight + 40) { el = pick(); window.__abScroller = el; }
  const doc = el === se;
  const r = doc ? {left: 0, top: 0, right: innerWidth, bottom: innerHeight} : el.getBoundingClientRect();
  return JSON.stringify({w: innerWidth, h: doc ? innerHeight : el.clientHeight, y: el.scrollTop, sh: el.scrollHeight,
    l: Math.max(r.left, 0), t: Math.max(r.top, 0), r: Math.min(r.right, innerWidth), b: Math.min(r.bottom, innerHeight)});
})()
"""


class UserMoved(Exception):
    pass


CAL_JS = r"""
(() => {
  if (!window.__abCalOn) {
    window.__abCalOn = true;
    document.addEventListener('mousemove', e => {
      if (e.isTrusted) window.__abCal = {dx: e.screenX - e.clientX, dy: e.screenY - e.clientY, t: Date.now()};
    }, true);
  }
  return JSON.stringify({sx: screenX, sy: screenY, ow: outerWidth, oh: outerHeight,
                         iw: innerWidth, ih: innerHeight, dpr: devicePixelRatio});
})()
"""


class RealMouse:
    """Moves and clicks with the actual Windows mouse pointer."""

    def __init__(self):
        import ctypes
        from ctypes import wintypes
        self.ct, self.wt = ctypes, wintypes
        u = ctypes.windll.user32
        u.WindowFromPoint.argtypes = [wintypes.POINT]
        u.WindowFromPoint.restype = wintypes.HWND
        u.GetAncestor.argtypes = [wintypes.HWND, wintypes.UINT]
        u.GetAncestor.restype = wintypes.HWND
        self.u = u
        self.last = None

    def pos(self):
        pt = self.wt.POINT()
        self.u.GetCursorPos(self.ct.byref(pt))
        return pt.x, pt.y

    def move(self, x, y):
        if self.last:
            cx, cy = self.pos()
            if abs(cx - self.last[0]) > 30 or abs(cy - self.last[1]) > 30:
                raise UserMoved()
        x, y = int(round(x)), int(round(y))
        self.u.SetCursorPos(x, y)
        self.last = (x, y)

    def chrome_at(self, x, y):
        h = self.u.WindowFromPoint(self.wt.POINT(int(x), int(y)))
        root = self.u.GetAncestor(h, 2)
        buf = self.ct.create_unicode_buffer(64)
        self.u.GetClassNameW(root, buf, 64)
        return buf.value.startswith("Chrome_WidgetWin")

    def button(self, down):
        self.u.mouse_event(0x0002 if down else 0x0004, 0, 0, 0, 0)

    def wheel(self, delta):
        self.u.mouse_event(0x0800, 0, 0, delta, 0)


HEALTH_JS = r"""
(() => {
  const nav = performance.getEntriesByType('navigation')[0];
  const pw = [...document.querySelectorAll('input[type="password"]')].some(e => {
    const r = e.getBoundingClientRect(); const st = getComputedStyle(e);
    return r.width > 20 && r.height > 8 && st.visibility !== 'hidden' && st.display !== 'none';
  });
  return JSON.stringify({url: location.href, status: (nav && nav.responseStatus) || 0, pw: pw,
                         title: document.title || '', text: (document.body ? document.body.innerText : '').slice(0, 2000)});
})()
"""

PAGE_TEXT_JS = r"""
(() => { const root = document.querySelector('main') || document.body; return root ? root.innerText.slice(0, 60000) : ''; })()
"""

PAGES_DIR = os.path.join(APP_DIR, "Page snapshots")
LOGIN_URL = re.compile(r"/(login|log-in|signin|sign-in|sign_in|logon|auth|session/new|account/login)\b", re.I)
DOWN_TEXT = re.compile(r"(bad gateway|service unavailable|gateway time-?out|internal server error|"
                       r"this site can.t be reached|err_[a-z_]+|site is down|under maintenance)", re.I)
TIME_BITS = re.compile(r"(\d{1,2}:\d{2}(:\d{2})?\s?(am|pm)?|\d+\s?(sec|second|min|minute|hour|hr|day)s?\s+ago|"
                       r"just now|yesterday|today)", re.I)


def page_lines(text):
    out = []
    for line in text.splitlines():
        line = TIME_BITS.sub("", line).strip()
        if len(line) > 2:
            out.append(line)
    return out


def compare_page(site, text):
    """Returns (changed, new_lines). The first visit only stores the page."""
    os.makedirs(PAGES_DIR, exist_ok=True)
    path = os.path.join(PAGES_DIR, site["id"] + ".txt")
    new = page_lines(text)
    try:
        with open(path, "r", encoding="utf-8") as f:
            old = f.read().splitlines()
    except OSError:
        old = None
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(new))
    if old is None or not new:
        return False, []
    old_set = set(old)
    added = [l for l in new if l not in old_set]
    ratio = difflib.SequenceMatcher(None, "\n".join(old), "\n".join(new), autojunk=False).ratio()
    return bool(added) and ratio < 0.995, added[:4]


class Visit:
    """One human-paced visit to a site inside a single Chrome tab."""

    def __init__(self, tab, site, stop_event):
        self.tab = tab
        self.site = site
        self.stop = stop_event
        self.name = host_of(site["url"]) or site["url"]
        self.real = RealMouse() if site.get("real_mouse") and os.name == "nt" else None
        self.cal = None

    # timing ---------------------------------------------------------------
    def pause(self, a, b):
        if self.stop.wait(random.uniform(a, b)):
            raise Stopped()

    def say(self, text):
        log.info("%s   %s", self.name, text)

    # page state -----------------------------------------------------------
    def view(self):
        try:
            return json.loads(self.tab.js(VIEW_JS) or "{}")
        except Exception:
            return {"w": 1200, "h": 700, "y": 0, "sh": 700}

    def mark(self):
        self.tab.js("window.__abMark = 1")

    def wait_load(self, old_href=None, grace=None, timeout=45):
        start = time.time()
        self.pause(0.7, 1.2)
        while time.time() - start < timeout:
            try:
                raw = self.tab.js("JSON.stringify([document.readyState, location.href, window.__abMark === 1])", 8)
                state, href, marked = json.loads(raw)
            except Exception:
                self.pause(0.4, 0.6)
                continue
            if state == "complete" and not href.startswith("about:"):
                if not marked or (old_href and href != old_href):
                    self.cal = None
                    return True
                if grace and time.time() - start > grace:
                    return False
            self.pause(0.4, 0.6)
        return False

    # mouse ----------------------------------------------------------------
    def calibrate(self):
        """Find where the page sits on screen, then confirm it with a real pointer movement."""
        g = json.loads(self.tab.js(CAL_JS))
        border = (g["ow"] - g["iw"]) / 2
        self.cal = (g["sx"] + border, g["sy"] + g["oh"] - g["ih"] - border, g["dpr"])
        tx, ty = g["iw"] * random.uniform(0.4, 0.6), g["ih"] * random.uniform(0.4, 0.6)
        for _ in range(3):
            sx, sy = self.to_screen(tx + random.uniform(-3, 3), ty + random.uniform(-3, 3))
            self.real.move(sx, sy)
            time.sleep(random.uniform(0.04, 0.09))
        time.sleep(0.15)
        raw = self.tab.js("JSON.stringify(window.__abCal || null)")
        c = json.loads(raw) if raw else None
        if c:
            self.cal = (c["dx"], c["dy"], g["dpr"])
        self.tab.mx, self.tab.my = tx, ty

    def to_screen(self, x, y):
        dx, dy, dpr = self.cal
        return (x + dx) * dpr, (y + dy) * dpr

    def mouse(self, kind, x, y, **extra):
        if self.real:
            if self.cal is None:
                self.calibrate()
            sx, sy = self.to_screen(x, y)
            if kind == "mouseMoved":
                self.real.move(sx, sy)
            elif kind == "mousePressed":
                if not self.real.chrome_at(sx, sy):
                    raise RuntimeError("Chrome is covered by another window, click skipped")
                self.real.move(sx, sy)
                self.real.button(True)
            elif kind == "mouseReleased":
                self.real.button(False)
            elif kind == "mouseWheel":
                if not self.real.chrome_at(sx, sy):
                    raise RuntimeError("Chrome is covered by another window, scroll skipped")
                self.real.wheel(-120 if extra.get("deltaY", 0) > 0 else 120)
            return
        p = {"type": kind, "x": round(x, 1), "y": round(y, 1), "pointerType": "mouse"}
        p.update(extra)
        self.tab.send("Input.dispatchMouseEvent", p)

    def move_to(self, x, y):
        sx, sy = self.tab.mx, self.tab.my
        dx, dy = x - sx, y - sy
        dist = math.hypot(dx, dy)
        if dist < 2:
            return
        spread = max(25.0, dist * 0.28)
        c1 = (sx + dx * random.uniform(0.15, 0.4) + random.uniform(-spread, spread) * 0.5,
              sy + dy * random.uniform(0.15, 0.4) + random.uniform(-spread, spread) * 0.5)
        c2 = (sx + dx * random.uniform(0.6, 0.85) + random.uniform(-spread, spread) * 0.3,
              sy + dy * random.uniform(0.6, 0.85) + random.uniform(-spread, spread) * 0.3)
        overshoot = dist > 220 and random.random() < 0.3
        ex, ey = x, y
        if overshoot:
            k = random.uniform(5, 14) / dist
            ex, ey = x + dx * k, y + dy * k
        steps = int(max(12, min(75, dist / random.uniform(9, 16))))
        for i in range(1, steps + 1):
            t = i / steps
            t = t * t * (3 - 2 * t)
            u = 1 - t
            px = u ** 3 * sx + 3 * u * u * t * c1[0] + 3 * u * t * t * c2[0] + t ** 3 * ex
            py = u ** 3 * sy + 3 * u * u * t * c1[1] + 3 * u * t * t * c2[1] + t ** 3 * ey
            self.mouse("mouseMoved", px, py)
            time.sleep(random.uniform(0.006, 0.02))
        if overshoot:
            time.sleep(random.uniform(0.05, 0.15))
            for i in range(1, 7):
                t = i / 6
                self.mouse("mouseMoved", ex + (x - ex) * t, ey + (y - ey) * t)
                time.sleep(random.uniform(0.01, 0.025))
        self.tab.mx, self.tab.my = x, y

    def drift(self):
        v = self.view()
        self.move_to(random.uniform(v["w"] * 0.15, v["w"] * 0.85), random.uniform(v["h"] * 0.2, v["h"] * 0.8))

    def click(self, x, y):
        x += random.uniform(-2, 2)
        y += random.uniform(-1.5, 1.5)
        self.move_to(x, y)
        self.pause(0.08, 0.35)
        self.mouse("mousePressed", x, y, button="left", buttons=1, clickCount=1)
        time.sleep(random.uniform(0.05, 0.14))
        self.mouse("mouseReleased", x, y, button="left", buttons=0, clickCount=1)

    # scrolling ------------------------------------------------------------
    def wheel(self, notches):
        """Scroll with mouse wheel notches (positive = down). Returns True if the page moved."""
        v = self.view()
        l, t, r, b = v.get("l", 0), v.get("t", 0), v.get("r", v["w"]), v.get("b", v["h"])
        if not (l + 10 < self.tab.mx < r - 10 and t + 10 < self.tab.my < b - 10):
            self.move_to(random.uniform(l + (r - l) * 0.3, l + (r - l) * 0.7),
                         random.uniform(t + (b - t) * 0.3, t + (b - t) * 0.7))
            self.pause(0.1, 0.4)
        before = v["y"]
        step = 100 if notches > 0 else -100
        for _ in range(abs(notches)):
            self.mouse("mouseWheel", self.tab.mx, self.tab.my, deltaX=0, deltaY=step)
            time.sleep(random.uniform(0.04, 0.12))
        time.sleep(0.35)
        after = self.view()
        if abs(after["y"] - before) > 4:
            return True
        at_bottom = after["y"] + after["h"] >= after["sh"] - 4
        if (notches > 0 and at_bottom) or (notches < 0 and after["y"] <= 0):
            return False
        self.tab.js("window.__abScroller && window.__abScroller.scrollBy({top: %d, behavior: document.hidden ? 'instant' : 'smooth'})"
                    % (step * abs(notches)))
        time.sleep(0.6)
        return abs(self.view()["y"] - before) > 4

    def read(self, seconds):
        end = time.time() + seconds
        scroll = self.site.get("scroll", True)
        while time.time() < end:
            r = random.random()
            v = self.view()
            at_bottom = v["y"] + v["h"] >= v["sh"] - 4
            if scroll and r < 0.55 and not at_bottom:
                self.wheel(random.randint(1, 4))
                self.pause(0.5, 2.4)
            elif scroll and r < 0.67 and v["y"] > 250:
                self.wheel(-random.randint(1, 2))
                self.pause(0.6, 1.6)
            elif r < 0.74:
                self.drift()
                self.pause(0.3, 1.3)
            elif r < 0.84:
                self.hover_something()
            else:
                self.idle(random.uniform(1.2, 3.8))

    def hover_something(self):
        """Rest the pointer on a link or button for a moment without clicking."""
        raw = self.tab.js(HOVER_JS)
        if not raw:
            self.idle(random.uniform(0.8, 2.0))
            return
        p = json.loads(raw)
        self.move_to(p["x"], p["y"])
        self.idle(random.uniform(0.4, 1.6))

    def idle(self, seconds):
        """Wait like a reader: hand mostly still, with tiny shifts now and then."""
        end = time.time() + seconds
        while time.time() < end:
            self.pause(0.4, 1.1)
            if random.random() < 0.35:
                self.mouse("mouseMoved", self.tab.mx + random.uniform(-4, 4), self.tab.my + random.uniform(-3, 3))

    # links ----------------------------------------------------------------
    def rect(self, idx):
        js = (RECT_JS.replace("INDEX", str(idx))
              .replace("FX", "%.2f" % random.uniform(0.3, 0.7))
              .replace("FY", "%.2f" % random.uniform(0.35, 0.65)))
        raw = self.tab.js(js)
        return json.loads(raw) if raw else None

    def bring_into_view(self, idx):
        for _ in range(40):
            info = self.rect(idx)
            if not info:
                return None
            vh = info["vh"]
            if vh * 0.12 < info["top"] and info["top"] + info["h"] < vh * 0.88:
                return info
            delta = info["top"] - vh * random.uniform(0.35, 0.55)
            notches = max(1, min(4, int(abs(delta) / 100)))
            moved = self.wheel(notches if delta > 0 else -notches)
            self.pause(0.15, 0.6)
            if not moved:
                self.tab.js("window.__abLinks[%d].scrollIntoView({block: 'center', behavior: 'smooth'})" % idx)
                self.pause(0.8, 1.4)
                return self.rect(idx)
        return self.rect(idx)

    def targets(self):
        return [w.strip() for w in self.site.get("targets", "").split(",") if w.strip()]

    def open_inner_page(self, visited):
        skip = [w.strip().lower() for w in self.site.get("skip", "").split(",") if w.strip()]
        raw = self.tab.js(LINKS_JS.replace("SKIP_LIST", json.dumps(skip)))
        links = json.loads(raw or "[]")
        wanted = self.targets()

        if wanted:
            # pick one of the user's chosen pages at random, then a matching link
            random.shuffle(wanted)
            pool, chosen = [], None
            for w in wanted:
                key = w.lower()
                pool = [l for l in links if key in (l["href"] + " " + l["text"]).lower()]
                if pool:
                    chosen = w
                    break
            if not pool:
                here = (self.tab.js("location.href") or "").split("#")[0]
                direct = []
                for w in wanted:
                    if w.startswith("/") or w.lower().startswith("http"):
                        url = urllib.parse.urljoin(self.site["url"], w)
                        if host_of(url) == host_of(self.site["url"]) and url != here:
                            direct.append((url, w))
                if direct:
                    url, w = random.choice(direct)
                    return self.navigate(url, w, visited)
                self.say("None of your chosen pages are linked here, going back")
                self.go_back()
                return False
            fresh = [l for l in pool if l["href"] not in visited] or pool
        else:
            fresh = [l for l in links if l["href"] not in visited] or links
            if not fresh:
                self.say("No inner links found on this page")
                return False

        random.shuffle(fresh)
        old = self.tab.js("location.href")
        for link in fresh[:4]:
            info = self.bring_into_view(link["i"])
            if not info or not info["ok"]:
                continue
            self.pause(0.3, 1.2)
            self.tab.js("(() => { const a = window.__abLinks[%d]; if (a && a.target === '_blank') a.target = '_self';"
                        " window.__abMark = 1; })()" % link["i"])
            self.click(info["cx"], info["cy"])
            if self.wait_load(old, grace=10):
                visited.add(link["href"])
                self.say("Clicked  %s" % (link["text"] or link["href"]))
                return True
        link = fresh[0]
        return self.navigate(link["href"], link["text"] or link["href"], visited)

    def navigate(self, url, label, visited):
        old = self.tab.js("location.href")
        self.mark()
        self.tab.send("Page.navigate", {"url": url})
        self.wait_load(old)
        visited.add(url)
        self.say("Opened  %s" % label)
        return True

    def go_back(self):
        old = self.tab.js("location.href")
        self.mark()
        self.tab.js("history.back()")
        if self.wait_load(old, grace=8):
            self.say("Went back")

    # whole visit ----------------------------------------------------------
    def run(self, refresh):
        s = self.site
        try:
            self.tab.send("Emulation.setFocusEmulationEnabled", {"enabled": True})
        except Exception:
            pass
        if refresh:
            self.mark()
            self.tab.send("Page.reload", {"ignoreCache": False})
            self.say("Refreshed")
        loaded = self.wait_load()
        result = {"status": "ok", "detail": ""}
        state = self.check_page(loaded)
        if state != "ok":
            result["status"] = state
            return result
        self.pause(1.5, 4.0)
        if s.get("alert_change", True):
            try:
                changed, added = compare_page(s, self.tab.js(PAGE_TEXT_JS) or "")
            except Exception:
                changed, added = False, []
            if changed:
                result["status"] = "changed"
                result["detail"] = "\n".join(added)
                self.say("The page changed")
        self.read(random.uniform(min(s["stay_min"], s["stay_max"]), max(s["stay_min"], s["stay_max"])))

        moved = False
        if s.get("browse"):
            lo, hi = sorted((int(s["pages_min"]), int(s["pages_max"])))
            visited = {s["url"].split("#")[0], s["url"].rstrip("/") + "/", s["url"].rstrip("/")}
            for n in range(random.randint(lo, hi)):
                if n and random.random() < 0.15:
                    self.go_back()
                    self.read(random.uniform(3, 8))
                if self.open_inner_page(visited):
                    moved = True
                    self.pause(0.8, 2.5)
                    self.read(random.uniform(min(s["stay_min"], s["stay_max"]), max(s["stay_min"], s["stay_max"])))

        if s.get("home") and moved:
            self.pause(1, 3)
            self.mark()
            self.tab.send("Page.navigate", {"url": s["url"]})
            self.wait_load()
            self.say("Back on the main page")
            self.pause(1, 2.5)
            self.read(random.uniform(3, 7))
        return result

    def check_page(self, loaded):
        """ok, down or signed_out."""
        try:
            h = json.loads(self.tab.js(HEALTH_JS, timeout=10) or "{}")
        except Exception:
            h = {}
        url = h.get("url", "")
        if not h or url.startswith("chrome-error") or (h.get("status") or 0) >= 500 or \
                (not loaded and DOWN_TEXT.search(h.get("title", "") + " " + h.get("text", "")[:400])):
            self.say("The site did not load properly")
            return "down"
        if h.get("pw") or (LOGIN_URL.search(urllib.parse.urlparse(url).path or "") and
                           not LOGIN_URL.search(urllib.parse.urlparse(self.site["url"]).path or "")):
            self.say("Signed out (login page showing)")
            return "signed_out"
        return "ok"


def chrome_running():
    try:
        if os.name == "nt":
            out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq chrome.exe", "/NH"], capture_output=True,
                                 text=True, creationflags=0x08000000, timeout=10).stdout
            return "chrome.exe" in out.lower()
        return subprocess.run(["pgrep", "-x", "chrome"], capture_output=True, timeout=10).returncode == 0
    except Exception:
        return False


def same_site(tab_url, host):
    h = host_of(tab_url)
    return bool(host) and (h == host or h.endswith("." + host))


def visit_site(cfg, site, stop_event):
    chrome = Chrome(cfg)
    url = site["url"]
    host = host_of(url)
    if not chrome.alive():
        BRIDGE.ready.wait(8)
    was_running = chrome.alive()
    if not was_running:
        if chrome_running():
            if not BRIDGE.ready.wait(20):
                raise RuntimeError("Chrome is open but the AutoBrowse Helper is not connected. "
                                   "Check it in Settings.")
        else:
            log.info("%s   Starting Chrome", host)
            chrome.launch()
        time.sleep(4)

    tabs = chrome.tabs()
    match = next((t for t in tabs if t.get("active") and same_site(t.get("url", ""), host)), None) \
        or next((t for t in tabs if same_site(t.get("url", ""), host)), None)

    refresh = True
    if not match:
        blank = next((t for t in tabs if t.get("url", "").startswith(("chrome://newtab", "about:blank",
                                                                       "chrome://new-tab-page"))), None)
        if blank:
            match = blank
            chrome.bridge.call("navigate", tabId=blank["id"], url=url)
            log.info("%s   Not open yet, opened in the empty tab", host)
            refresh = False
            for _ in range(20):
                time.sleep(0.5)
                cur = next((t for t in chrome.tabs() if t["id"] == blank["id"]), None)
                if cur and same_site(cur.get("url", ""), host):
                    break
        else:
            match = chrome.new_tab(url, bool(site.get("front") or site.get("real_mouse")))
            log.info("%s   Not open yet, opened in a new tab", host)
            refresh = False
            time.sleep(1)

    if site.get("real_mouse") and os.name == "nt":
        chrome.focus(match["id"])
        time.sleep(0.8)
    elif site.get("front"):
        chrome.activate(match["id"])

    tab = Tab(match)
    try:
        return Visit(tab, site, stop_event).run(refresh)
    finally:
        tab.close()


# ----------------------------------------------------------------------------
# Active time
# ----------------------------------------------------------------------------

STATS_PATH = os.path.join(APP_DIR, "active_time.json")


class ActiveTime:
    """Remembers how long AutoBrowse spent visiting websites."""

    def __init__(self):
        self.lock = threading.Lock()
        self.visits = []
        self.alerts = []
        self.reset_at = time.time()
        try:
            with open(STATS_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
            self.visits = data.get("visits", [])
            self.alerts = data.get("alerts", [])
            self.reset_at = data.get("reset_at", self.reset_at)
        except (OSError, ValueError):
            pass

    def save(self):
        os.makedirs(APP_DIR, exist_ok=True)
        cutoff = time.time() - 60 * 86400
        self.visits = [v for v in self.visits if v[0] >= cutoff]
        self.alerts = [a for a in self.alerts if a[0] >= cutoff]
        with open(STATS_PATH, "w", encoding="utf-8") as f:
            json.dump({"visits": self.visits, "alerts": self.alerts, "reset_at": self.reset_at}, f)

    def add(self, start, seconds, host, site_id="", result="ok"):
        with self.lock:
            self.visits.append([start, round(seconds, 1), host, site_id, result])
            self.save()

    def alert(self, host, kind, detail=""):
        with self.lock:
            self.alerts.append([time.time(), host, kind, detail])
            self.save()

    def today_for(self, site_id):
        start = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0).timestamp()
        with self.lock:
            rows = [v for v in self.visits if v[0] >= start and len(v) > 3 and v[3] == site_id]
        return sum(v[1] for v in rows), len(rows)

    def reset(self):
        with self.lock:
            self.reset_at = time.time()
            self.save()

    def total(self, since):
        with self.lock:
            rows = [v for v in self.visits if v[0] >= since]
        return sum(v[1] for v in rows), len(rows)


ACTIVE = ActiveTime()


def idle_seconds():
    """Seconds since the keyboard or mouse was last used (Windows)."""
    if os.name != "nt":
        return 1e9
    import ctypes

    class LASTINPUTINFO(ctypes.Structure):
        _fields_ = [("cbSize", ctypes.c_uint), ("dwTime", ctypes.c_uint)]
    info = LASTINPUTINFO()
    info.cbSize = ctypes.sizeof(info)
    if not ctypes.windll.user32.GetLastInputInfo(ctypes.byref(info)):
        return 1e9
    return ((ctypes.windll.kernel32.GetTickCount() - info.dwTime) & 0xFFFFFFFF) / 1000.0


def on_battery():
    if os.name != "nt":
        return False
    import ctypes

    class POWER(ctypes.Structure):
        _fields_ = [("ACLineStatus", ctypes.c_ubyte), ("BatteryFlag", ctypes.c_ubyte),
                    ("BatteryLifePercent", ctypes.c_ubyte), ("SystemStatusFlag", ctypes.c_ubyte),
                    ("BatteryLifeTime", ctypes.c_ulong), ("BatteryFullLifeTime", ctypes.c_ulong)]
    st = POWER()
    if not ctypes.windll.kernel32.GetSystemPowerStatus(ctypes.byref(st)):
        return False
    return st.ACLineStatus == 0


def span(seconds):
    m = int(seconds // 60)
    return "%dh %02dm" % (m // 60, m % 60) if m >= 60 else "%dm %02ds" % (m, int(seconds % 60))


# ----------------------------------------------------------------------------
# Scheduler
# ----------------------------------------------------------------------------

class Scheduler(threading.Thread):
    def __init__(self, app):
        super().__init__(daemon=True)
        self.app = app
        self.next_run = {}
        self.forced = set()
        self.running_id = None
        self.paused = False
        self.stop_event = threading.Event()
        self.lock = threading.Lock()
        self.blocked = {}
        self.wait_reason = ""
        self.site_state = {}

    def plan(self, site, first=False):
        if first:
            delay = random.uniform(8, 25)
        elif site["mode"] == "fixed":
            delay = float(site["fixed_min"]) * 60
        else:
            lo, hi = sorted((float(site["rand_min"]), float(site["rand_max"])))
            delay = random.uniform(lo, hi) * 60
        self.next_run[site["id"]] = time.time() + max(10.0, delay)

    def sync(self, sites):
        with self.lock:
            ids = {s["id"] for s in sites}
            for s in sites:
                if s["id"] not in self.next_run:
                    self.plan(s, first=True)
            for k in list(self.next_run):
                if k not in ids:
                    self.next_run.pop(k, None)

    def replan(self, site):
        with self.lock:
            self.plan(site)

    def run_now(self, site_id=None):
        with self.lock:
            if site_id:
                self.forced.add(site_id)
            else:
                self.forced.update(s["id"] for s in self.app.cfg["sites"] if s.get("enabled"))

    def run(self):
        while not self.stop_event.is_set():
            cfg = self.app.cfg
            sites = [s for s in cfg["sites"] if s.get("enabled") and s.get("url")]
            self.sync(sites)
            now = time.time()
            allowed = not self.paused and within_hours(cfg)
            self.wait_reason = ""
            if allowed and cfg.get("pause_busy") and idle_seconds() < float(cfg.get("busy_minutes") or 2) * 60:
                allowed, self.wait_reason = False, "Waiting: you are using the PC"
            elif allowed and cfg.get("pause_battery") and on_battery():
                allowed, self.wait_reason = False, "Waiting: on battery"
            blocked = {}
            for s in sites:
                if s.get("hours_on") and not within_hours(s):
                    blocked[s["id"]] = "Outside its hours (%s to %s)" % (s.get("hours_from"), s.get("hours_to"))
                elif s.get("limit_on"):
                    secs, n = ACTIVE.today_for(s["id"])
                    if (int(s.get("limit_visits") or 0) and n >= int(s["limit_visits"])) or \
                            (float(s.get("limit_minutes") or 0) and secs >= float(s["limit_minutes"]) * 60):
                        blocked[s["id"]] = "Daily limit reached"
            self.blocked = blocked
            with self.lock:
                due = [s for s in sites
                       if s["id"] in self.forced or
                       (allowed and s["id"] not in blocked and self.next_run.get(s["id"], now + 1) <= now)]
            due.sort(key=lambda s: self.next_run.get(s["id"], 0))
            for s in due:
                if self.stop_event.is_set():
                    break
                with self.lock:
                    self.forced.discard(s["id"])
                self.running_id = s["id"]
                started = time.time()
                outcome = "error"
                try:
                    res = visit_site(cfg, s, self.stop_event) or {"status": "ok", "detail": ""}
                    outcome = res["status"]
                    self.handle_result(s, res)
                except Stopped:
                    outcome = "stopped"
                    break
                except UserMoved:
                    log.info("%s   You moved the mouse, so this visit stopped", host_of(s["url"]))
                except Exception as e:
                    log.info("%s   Problem: %s", host_of(s["url"]), e)
                finally:
                    ACTIVE.add(started, time.time() - started, host_of(s["url"]), s["id"], outcome)
                    self.running_id = None
                    self.replan(s)
            self.stop_event.wait(1)

    def handle_result(self, s, res):
        host = host_of(s["url"])
        st = self.site_state.setdefault(s["id"], {"down": False, "signed_out": False})
        status = res["status"]
        if status == "down":
            if not st["down"] and s.get("alert_down", True):
                ACTIVE.alert(host, "Site down")
                self.app.alert("%s is down" % host, "The page did not load or showed an error. "
                                                    "AutoBrowse will keep checking and tell you when it is back.")
            st["down"] = True
            return
        if st["down"]:
            st["down"] = False
            if s.get("alert_down", True):
                ACTIVE.alert(host, "Site back up")
                self.app.alert("%s is back up" % host, "The page is loading normally again.")
        if status == "signed_out":
            if not st["signed_out"] and s.get("alert_logout", True):
                ACTIVE.alert(host, "Signed out")
                self.app.alert("Signed out of %s" % host, "The login page is showing. Sign in again in Chrome "
                                                         "so AutoBrowse can keep visiting.")
            st["signed_out"] = True
            return
        st["signed_out"] = False
        if status == "changed" and s.get("alert_change", True):
            ACTIVE.alert(host, "Page changed", res.get("detail", ""))
            body = res.get("detail") or "Something new appeared on the page."
            self.app.alert("%s changed" % host, body[:300])


# ----------------------------------------------------------------------------
# Tray icon
# ----------------------------------------------------------------------------

def make_icon(size=64, paused=False):
    big = size * 4
    img = Image.new("RGBA", (big, big), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    top, bottom = ((150, 156, 161), (95, 99, 104)) if paused else ((52, 211, 120), (21, 128, 61))
    for y in range(big):
        t = y / big
        col = tuple(int(top[i] + (bottom[i] - top[i]) * t) for i in range(3))
        d.line([(0, y), (big, y)], fill=col)
    mask = Image.new("L", (big, big), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, big - 1, big - 1), radius=int(big * 0.24), fill=255)
    img.putalpha(mask)
    c = big / 2
    r = big * 0.25
    w = int(big * 0.085)
    d.arc((c - r, c - r, c + r, c + r), start=60, end=330, fill="white", width=w)
    ang = math.radians(330)
    px, py = c + (r - w / 2) * math.cos(ang), c + (r - w / 2) * math.sin(ang)
    tx, ty = -math.sin(ang), math.cos(ang)
    nx, ny = math.cos(ang), math.sin(ang)
    a, b = big * 0.13, big * 0.12
    d.polygon([(px + tx * a, py + ty * a), (px + nx * b, py + ny * b), (px - nx * b, py - ny * b)], fill="white")
    dot = big * 0.055
    d.ellipse((c - dot, c - dot, c + dot, c + dot), fill="white")
    return img.resize((size, size), Image.LANCZOS)


def resource(name):
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)


# ----------------------------------------------------------------------------
# Main window
# ----------------------------------------------------------------------------

GREEN = "#16a34a"
GREEN_HOVER = "#15803d"
SIDEBAR = "#0d2818"
SIDEBAR_HOVER = "#173d27"
PAGE_BG = ("#f3f6f4", "#0f1411")
CARD_BG = ("#ffffff", "#182019")
CARD_BORDER = ("#e2e8e4", "#26322a")
MUTED = ("#64748b", "#94a3b8")
TEXT = ("#0f172a", "#e8efe9")
RED = "#dc2626"


class SmoothScrollFrame(ctk.CTkScrollableFrame):
    """Scrollable panel with eased, animated mouse wheel scrolling."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._target = None
        self._anim = None

    def _sizes(self):
        c = self._parent_canvas
        box = c.bbox("all")
        if not box:
            return None
        total = box[3] - box[1]
        view = c.winfo_height()
        return (total, view) if total > view else None

    def _mouse_wheel_all(self, event):
        if not self._check_if_valid_scroll(event.widget) or self._shift_pressed:
            return
        sizes = self._sizes()
        if not sizes:
            return
        total, view = sizes
        if sys.platform.startswith("win") or sys.platform == "darwin":
            notches = event.delta / 120.0
        else:
            notches = 1 if getattr(event, "num", 5) == 4 else -1
        current = self._parent_canvas.yview()[0] * total
        base = self._target if self._target is not None else current
        self._target = min(max(base - notches * 110, 0), total - view)
        if self._anim is None:
            self._step()

    def _step(self):
        sizes = self._sizes()
        if not sizes or self._target is None:
            self._anim, self._target = None, None
            return
        total, _ = sizes
        current = self._parent_canvas.yview()[0] * total
        diff = self._target - current
        if abs(diff) < 1.5:
            self._parent_canvas.yview_moveto(self._target / total)
            self._anim, self._target = None, None
            return
        self._parent_canvas.yview_moveto((current + diff * 0.22) / total)
        self._anim = self.after(10, self._step)


def font(size=13, weight="normal"):
    return ctk.CTkFont(family="Segoe UI", size=size, weight=weight)


REPORT_DIR = os.path.join(os.path.expanduser("~"), "Documents", "AutoBrowse Reports")


def export_report(days=30):
    """Write a formatted Excel report of the last `days` days and return its path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    since = time.time() - days * 86400
    with ACTIVE.lock:
        visits = [v for v in ACTIVE.visits if v[0] >= since]
        alerts = [a for a in ACTIVE.alerts if a[0] >= since]

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="166534")
    head_font = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Segoe UI", size=10)
    band = PatternFill("solid", fgColor="F0F7F2")
    thin = Side(style="thin", color="D9E3DC")
    border = Border(bottom=thin)

    def sheet(ws, title, headers, rows, widths, formats):
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.row_dimensions[1].height = 24
        for r in rows:
            ws.append(r)
        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            for j, c in enumerate(row):
                c.font, c.border = body_font, border
                if formats.get(j):
                    c.number_format = formats[j]
                if i % 2 == 0:
                    c.fill = band
                c.alignment = Alignment(vertical="top", wrap_text=(headers[j] in ("Details", "Website")))
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(rows) + 1)

    summary = {}
    for v in visits:
        d = datetime.fromtimestamp(v[0]).date()
        key = (d, v[2])
        row = summary.setdefault(key, {"visits": 0, "secs": 0.0, "changed": 0, "out": 0, "down": 0})
        row["visits"] += 1
        row["secs"] += v[1]
        res = v[4] if len(v) > 4 else "ok"
        row["changed"] += res == "changed"
        row["out"] += res == "signed_out"
        row["down"] += res == "down"
    srows = [[d, host, r["visits"], round(r["secs"] / 60, 1), r["changed"], r["out"], r["down"]]
             for (d, host), r in sorted(summary.items(), key=lambda kv: (kv[0][0], kv[0][1]), reverse=True)]
    names = {"ok": "Visited", "changed": "Page changed", "signed_out": "Signed out", "down": "Site down",
             "error": "Problem", "stopped": "Stopped"}
    vrows = [[datetime.fromtimestamp(v[0]).date(), datetime.fromtimestamp(v[0]).time().replace(microsecond=0),
              v[2], round(v[1] / 60, 2), names.get(v[4] if len(v) > 4 else "ok", "Visited")]
             for v in sorted(visits, key=lambda v: v[0], reverse=True)]
    arows = [[datetime.fromtimestamp(a[0]).date(), datetime.fromtimestamp(a[0]).time().replace(microsecond=0),
              a[1], a[2], a[3]] for a in sorted(alerts, key=lambda a: a[0], reverse=True)]

    sheet(wb.active, "Summary", ["Date", "Website", "Visits", "Active minutes", "Page changes", "Signed out",
                                 "Site down"], srows, [13, 30, 9, 15, 14, 12, 11],
          {0: "dd mmm yyyy", 3: "0.0"})
    sheet(wb.create_sheet(), "Visits", ["Date", "Time", "Website", "Minutes", "Result"], vrows,
          [13, 11, 30, 10, 16], {0: "dd mmm yyyy", 1: "hh:mm:ss AM/PM", 3: "0.00"})
    sheet(wb.create_sheet(), "Alerts", ["Date", "Time", "Website", "Alert", "Details"], arows,
          [13, 11, 28, 16, 60], {0: "dd mmm yyyy", 1: "hh:mm:ss AM/PM"})
    if srows:
        ws = wb["Summary"]
        n = len(srows) + 1
        ws.append([])
        ws.append(["Total", "", "=SUM(C2:C%d)" % n, "=SUM(D2:D%d)" % n, "=SUM(E2:E%d)" % n,
                   "=SUM(F2:F%d)" % n, "=SUM(G2:G%d)" % n])
        for c in ws[n + 2]:
            c.font = Font(name="Segoe UI", bold=True, size=10)
            c.border = Border(top=Side(style="medium", color="166534"))
        ws.cell(row=n + 2, column=4).number_format = "0.0"

    os.makedirs(REPORT_DIR, exist_ok=True)
    path = os.path.join(REPORT_DIR, "AutoBrowse Report %s.xlsx" % datetime.now().strftime("%Y-%m-%d %H%M"))
    wb.save(path)
    return path


class App:
    def __init__(self, server_socket):
        self.cfg = load_config()
        self.events = queue.Queue()
        self.server = server_socket
        self.sel_id = None
        self.hint_shown = False
        self.cards = {}

        ctk.set_appearance_mode(self.cfg.get("theme", "System"))
        ctk.set_default_color_theme("green")
        self.root = ctk.CTk(fg_color=PAGE_BG)
        self.root.title(APP_NAME)
        self.root.geometry("1120x720")
        self.root.minsize(1000, 640)
        self.root.protocol("WM_DELETE_WINDOW", self.hide)
        self.icon_photo = ImageTk.PhotoImage(make_icon(64))
        self.root.iconphoto(True, self.icon_photo)
        if os.name == "nt" and os.path.exists(resource("AutoBrowse.ico")):
            self.root.after(300, lambda: self.root.iconbitmap(resource("AutoBrowse.ico")))

        BRIDGE.start()
        try:
            write_helper()
        except OSError:
            pass
        self.scheduler = Scheduler(self)

        self.build_ui()
        self.scheduler.start()

        self.tray = pystray.Icon(APP_NAME, make_icon(64), APP_NAME, menu=pystray.Menu(
            pystray.MenuItem("Open AutoBrowse", lambda: self.events.put("show"), default=True),
            pystray.MenuItem("Visit all now", lambda: self.events.put("runall")),
            pystray.MenuItem(lambda item: "Resume" if self.scheduler.paused else "Pause",
                             lambda: self.events.put("toggle")),
            pystray.Menu.SEPARATOR,
            pystray.MenuItem("Exit", lambda: self.events.put("exit")),
        ))
        threading.Thread(target=self.tray.run, daemon=True).start()
        threading.Thread(target=self.listen, daemon=True).start()

        try:
            set_startup(self.cfg.get("startup", True))
        except OSError:
            pass
        threading.Thread(target=make_shortcuts, daemon=True).start()
        self.toast = None

        self.root.after(200, self.poll)
        self.root.after(1000, self.tick)
        log.info("AutoBrowse started")

    # ------------------------------------------------------------------ layout
    def build_ui(self):
        self.root.grid_columnconfigure(1, weight=1)
        self.root.grid_rowconfigure(0, weight=1)

        side = ctk.CTkFrame(self.root, width=230, corner_radius=0, fg_color=SIDEBAR)
        side.grid(row=0, column=0, sticky="nsw")
        side.grid_propagate(False)
        side.grid_columnconfigure(0, weight=1)
        side.grid_rowconfigure(5, weight=1)

        brand = ctk.CTkFrame(side, fg_color="transparent")
        brand.grid(row=0, column=0, sticky="ew", padx=20, pady=(26, 30))
        self.logo = ctk.CTkImage(light_image=make_icon(96), dark_image=make_icon(96), size=(42, 42))
        ctk.CTkLabel(brand, image=self.logo, text="").pack(side="left")
        names = ctk.CTkFrame(brand, fg_color="transparent")
        names.pack(side="left", padx=12)
        ctk.CTkLabel(names, text="AutoBrowse", font=font(20, "bold"), text_color="white").pack(anchor="w")
        ctk.CTkLabel(names, text="Smart page refresher", font=font(12), text_color="#86efac").pack(anchor="w")

        self.nav = {}
        for i, (key, label) in enumerate((("sites", "Websites"), ("settings", "Settings"), ("activity", "Activity"))):
            b = ctk.CTkButton(side, text=label, anchor="w", height=42, corner_radius=10, font=font(14, "bold"),
                              fg_color="transparent", hover_color=SIDEBAR_HOVER, text_color="#d1fae5",
                              command=lambda k=key: self.show_page(k))
            b.grid(row=1 + i, column=0, sticky="ew", padx=14, pady=3)
            self.nav[key] = b

        status = ctk.CTkFrame(side, fg_color=SIDEBAR_HOVER, corner_radius=14)
        status.grid(row=6, column=0, sticky="ew", padx=14, pady=(0, 10))
        top = ctk.CTkFrame(status, fg_color="transparent")
        top.pack(fill="x", padx=14, pady=(12, 0))
        self.status_dot = ctk.CTkLabel(top, text="●", font=font(16), text_color="#4ade80")
        self.status_dot.pack(side="left")
        self.status_text = ctk.CTkLabel(top, text="Running", font=font(14, "bold"), text_color="white")
        self.status_text.pack(side="left", padx=6)
        self.status_sub = ctk.CTkLabel(status, text="", font=font(12), text_color="#a7f3d0")
        self.status_sub.pack(anchor="w", padx=14, pady=(0, 10))
        self.pause_btn = ctk.CTkButton(status, text="Pause", height=34, corner_radius=8, font=font(13, "bold"),
                                       fg_color="transparent", border_width=1, border_color="#4ade80",
                                       hover_color=SIDEBAR, text_color="white", command=self.toggle_pause)
        self.pause_btn.pack(fill="x", padx=14, pady=(0, 14))
        ctk.CTkButton(side, text="Hide to tray", height=34, corner_radius=8, font=font(13),
                      fg_color="transparent", hover_color=SIDEBAR_HOVER, text_color="#a7f3d0",
                      command=self.hide).grid(row=7, column=0, sticky="ew", padx=14, pady=(0, 6))
        ctk.CTkLabel(side, text="Developed by Sourab Kumar Saha", font=font(11),
                     text_color="#86efac").grid(row=8, column=0, sticky="ew", padx=14, pady=(0, 14))

        self.pages = {}
        self.pages["sites"] = self.build_sites_page()
        self.pages["settings"] = self.build_settings_page()
        self.pages["activity"] = self.build_activity_page()
        self.show_page("sites")

    def show_page(self, key):
        for k, page in self.pages.items():
            if k == key:
                page.grid(row=0, column=1, sticky="nsew", padx=28, pady=24)
            else:
                page.grid_forget()
        for k, b in self.nav.items():
            b.configure(fg_color=GREEN if k == key else "transparent",
                        text_color="white" if k == key else "#d1fae5")

    def header(self, parent, title, subtitle):
        h = ctk.CTkFrame(parent, fg_color="transparent")
        h.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 18))
        ctk.CTkLabel(h, text=title, font=font(26, "bold"), text_color=TEXT).pack(anchor="w")
        ctk.CTkLabel(h, text=subtitle, font=font(13), text_color=MUTED).pack(anchor="w")
        return h

    def card(self, parent, title=None):
        c = ctk.CTkFrame(parent, fg_color=CARD_BG, corner_radius=14, border_width=1, border_color=CARD_BORDER)
        if title:
            ctk.CTkLabel(c, text=title, font=font(15, "bold"), text_color=TEXT).pack(anchor="w", padx=18, pady=(14, 6))
        return c

    def row(self, parent, pady=(4, 4)):
        r = ctk.CTkFrame(parent, fg_color="transparent")
        r.pack(fill="x", padx=18, pady=pady)
        return r

    def small_entry(self, parent, var, width=64):
        e = ctk.CTkEntry(parent, textvariable=var, width=width, height=32, corner_radius=8, justify="center")
        e.pack(side="left", padx=6)
        return e

    def label(self, parent, text, muted=False, side="left"):
        ctk.CTkLabel(parent, text=text, font=font(13), text_color=MUTED if muted else TEXT).pack(side=side)

    # ------------------------------------------------------------------ websites page
    def build_sites_page(self):
        page = ctk.CTkFrame(self.root, fg_color="transparent")
        page.grid_columnconfigure(1, weight=1)
        page.grid_rowconfigure(1, weight=1)
        h = self.header(page, "Websites", "Each website is refreshed and browsed on its own schedule.")
        ctk.CTkButton(h, text="+  Add website", height=40, corner_radius=10, font=font(14, "bold"),
                      fg_color=GREEN, hover_color=GREEN_HOVER, command=self.add_site).place(relx=1.0, rely=0.5, anchor="e")

        self.list_frame = SmoothScrollFrame(page, width=310, fg_color="transparent")
        self.list_frame.grid(row=1, column=0, sticky="nsw", padx=(0, 18))

        self.editor_wrap = ctk.CTkFrame(page, fg_color="transparent")
        self.editor_wrap.grid(row=1, column=1, sticky="nsew")
        self.editor_wrap.grid_columnconfigure(0, weight=1)
        self.editor_wrap.grid_rowconfigure(0, weight=1)

        self.empty = self.card(self.editor_wrap)
        ctk.CTkLabel(self.empty, image=ctk.CTkImage(make_icon(128), make_icon(128), size=(72, 72)), text="").pack(pady=(70, 10))
        ctk.CTkLabel(self.empty, text="No website selected", font=font(18, "bold"), text_color=TEXT).pack()
        ctk.CTkLabel(self.empty, text="Pick a website on the left, or click Add website to start.",
                     font=font(13), text_color=MUTED).pack(pady=(4, 70))

        ed = SmoothScrollFrame(self.editor_wrap, fg_color="transparent")
        self.editor = ed

        self.v_url = tk.StringVar()
        self.v_enabled = tk.IntVar(value=1)
        self.v_mode = tk.StringVar(value="Random")
        self.v_fixed = tk.StringVar()
        self.v_rmin = tk.StringVar()
        self.v_rmax = tk.StringVar()
        self.v_browse = tk.IntVar()
        self.v_pmin = tk.StringVar()
        self.v_pmax = tk.StringVar()
        self.v_smin = tk.StringVar()
        self.v_smax = tk.StringVar()
        self.v_scroll = tk.IntVar()
        self.v_home = tk.IntVar()
        self.v_front = tk.IntVar()
        self.v_real = tk.IntVar()
        self.v_al_change = tk.IntVar()
        self.v_al_logout = tk.IntVar()
        self.v_al_down = tk.IntVar()
        self.v_sh_on = tk.IntVar()
        self.v_sh_from = tk.StringVar()
        self.v_sh_to = tk.StringVar()
        self.v_lim_on = tk.IntVar()
        self.v_lim_visits = tk.StringVar()
        self.v_lim_min = tk.StringVar()

        c = self.card(ed, "Website")
        c.pack(fill="x", pady=(0, 12))
        self.url_entry = ctk.CTkEntry(c, textvariable=self.v_url, height=38, corner_radius=8, font=font(14))
        self.url_entry.pack(fill="x", padx=18, pady=(2, 10))
        ctk.CTkSwitch(c, text="Active", variable=self.v_enabled, font=font(13), progress_color=GREEN).pack(
            anchor="w", padx=18, pady=(0, 16))

        c = self.card(ed, "Refresh timing")
        c.pack(fill="x", pady=(0, 12))
        ctk.CTkSegmentedButton(c, values=["Fixed", "Random"], variable=self.v_mode, font=font(13, "bold"),
                               selected_color=GREEN, selected_hover_color=GREEN_HOVER,
                               command=lambda _v: self.timing_rows()).pack(anchor="w", padx=18, pady=(2, 8))
        self.fixed_row = ctk.CTkFrame(c, fg_color="transparent")
        self.label(self.fixed_row, "Refresh every")
        self.small_entry(self.fixed_row, self.v_fixed)
        self.label(self.fixed_row, "minutes")
        self.rand_row = ctk.CTkFrame(c, fg_color="transparent")
        self.label(self.rand_row, "Randomly between")
        self.small_entry(self.rand_row, self.v_rmin)
        self.label(self.rand_row, "and")
        self.small_entry(self.rand_row, self.v_rmax)
        self.label(self.rand_row, "minutes")
        self.timing_spacer = ctk.CTkFrame(c, fg_color="transparent", height=8)

        c = self.card(ed, "Browsing behavior")
        c.pack(fill="x", pady=(0, 12))
        ctk.CTkSwitch(c, text="Click through inner pages after each refresh", variable=self.v_browse,
                      font=font(13), progress_color=GREEN).pack(anchor="w", padx=18, pady=(2, 8))
        r = self.row(c)
        self.label(r, "Pages per visit")
        self.small_entry(r, self.v_pmin)
        self.label(r, "to")
        self.small_entry(r, self.v_pmax)
        r = self.row(c)
        self.label(r, "Time on each page")
        self.small_entry(r, self.v_smin)
        self.label(r, "to")
        self.small_entry(r, self.v_smax)
        self.label(r, "seconds")
        for text, var in (("Scroll and read pages naturally", self.v_scroll),
                          ("Return to the main page when done", self.v_home),
                          ("Switch to the tab while browsing (recommended)", self.v_front)):
            ctk.CTkSwitch(c, text=text, variable=var, font=font(13), progress_color=GREEN).pack(
                anchor="w", padx=18, pady=(8, 0))
        ctk.CTkSwitch(c, text="Use my real mouse pointer", variable=self.v_real, font=font(13, "bold"),
                      progress_color=GREEN).pack(anchor="w", padx=18, pady=(14, 0))
        ctk.CTkLabel(c, text="Moves your actual cursor and clicks with it, like a hand on the mouse. Chrome comes "
                             "to the front during the visit. Moving the mouse yourself stops the visit.",
                     font=font(12), text_color=MUTED, justify="left", anchor="w", wraplength=380).pack(
            anchor="w", padx=(64, 18))
        ctk.CTkFrame(c, fg_color="transparent", height=14).pack()

        c = self.card(ed, "Alerts")
        c.pack(fill="x", pady=(0, 12))
        for text, var in (("Alert me when the page changes", self.v_al_change),
                          ("Alert me when I'm signed out", self.v_al_logout),
                          ("Alert me when the site is down", self.v_al_down)):
            ctk.CTkSwitch(c, text=text, variable=var, font=font(13), progress_color=GREEN).pack(
                anchor="w", padx=18, pady=(2, 6))
        ctk.CTkFrame(c, fg_color="transparent", height=8).pack()

        c = self.card(ed, "Limits")
        c.pack(fill="x", pady=(0, 12))
        r = self.row(c)
        ctk.CTkSwitch(r, text="Only visit between", variable=self.v_sh_on, font=font(13),
                      progress_color=GREEN).pack(side="left")
        self.small_entry(r, self.v_sh_from, 70)
        self.label(r, "and")
        self.small_entry(r, self.v_sh_to, 70)
        r = self.row(c, (8, 4))
        ctk.CTkSwitch(r, text="Daily limit:", variable=self.v_lim_on, font=font(13),
                      progress_color=GREEN).pack(side="left")
        self.small_entry(r, self.v_lim_visits, 56)
        self.label(r, "visits and")
        self.small_entry(r, self.v_lim_min, 56)
        self.label(r, "minutes")
        ctk.CTkLabel(c, text="Times use 24 hour format. Use 0 for no limit on visits or minutes.",
                     font=font(12), text_color=MUTED).pack(anchor="w", padx=18, pady=(2, 14))

        c = self.card(ed, "Pages to click")
        c.pack(fill="x", pady=(0, 12))
        ctk.CTkLabel(c, text="Link names or page addresses separated by commas, for example Orders, Reports, "
                             "/dashboard. One is picked at random each time. Leave empty to click any page.",
                     font=font(12), text_color=MUTED, justify="left", anchor="w", wraplength=400).pack(
            anchor="w", padx=18)
        self.targets_box = ctk.CTkTextbox(c, height=64, corner_radius=8, font=font(13), wrap="word")
        self.targets_box.pack(fill="x", padx=18, pady=(6, 16))

        c = self.card(ed, "Never click links containing")
        c.pack(fill="x", pady=(0, 12))
        ctk.CTkLabel(c, text="Separate words with commas. Matched against the link text and address.",
                     font=font(12), text_color=MUTED, justify="left", anchor="w", wraplength=400).pack(
            anchor="w", padx=18)
        self.skip_box = ctk.CTkTextbox(c, height=74, corner_radius=8, font=font(13), wrap="word")
        self.skip_box.pack(fill="x", padx=18, pady=(6, 16))

        actions = ctk.CTkFrame(ed, fg_color="transparent")
        actions.pack(fill="x", pady=(2, 10))
        ctk.CTkButton(actions, text="Save", width=120, height=40, corner_radius=10, font=font(14, "bold"),
                      fg_color=GREEN, hover_color=GREEN_HOVER, command=self.save).pack(side="left")
        ctk.CTkButton(actions, text="Visit now", width=120, height=40, corner_radius=10, font=font(14, "bold"),
                      fg_color="transparent", border_width=2, border_color=GREEN, text_color=(GREEN, "#4ade80"),
                      hover_color=("#dcfce7", "#173d27"), command=self.visit_now).pack(side="left", padx=10)
        ctk.CTkButton(actions, text="Remove", width=100, height=40, corner_radius=10, font=font(14),
                      fg_color="transparent", text_color=RED, hover_color=("#fee2e2", "#3b1515"),
                      command=self.remove_site).pack(side="left")
        self.saved_label = ctk.CTkLabel(actions, text="", font=font(13, "bold"), text_color=GREEN)
        self.saved_label.pack(side="left", padx=12)

        self.render_cards()
        if self.cfg["sites"]:
            self.select(self.cfg["sites"][0]["id"])
        else:
            self.show_editor(False)
        return page

    def timing_rows(self):
        self.fixed_row.pack_forget()
        self.rand_row.pack_forget()
        self.timing_spacer.pack_forget()
        (self.fixed_row if self.v_mode.get() == "Fixed" else self.rand_row).pack(fill="x", padx=18, pady=(4, 4))
        self.timing_spacer.pack()

    def show_editor(self, on):
        if on:
            self.empty.grid_forget()
            self.editor.grid(row=0, column=0, sticky="nsew")
        else:
            self.editor.grid_forget()
            self.empty.grid(row=0, column=0, sticky="nsew")

    def render_cards(self):
        for w in self.list_frame.winfo_children():
            w.destroy()
        self.cards = {}
        if not self.cfg["sites"]:
            ctk.CTkLabel(self.list_frame, text="No websites yet", font=font(13), text_color=MUTED).pack(pady=30)
            return
        for s in self.cfg["sites"]:
            f = ctk.CTkFrame(self.list_frame, fg_color=CARD_BG, corner_radius=14, border_width=2,
                             border_color=CARD_BORDER, cursor="hand2")
            f.pack(fill="x", pady=5, padx=(0, 6))
            top = ctk.CTkFrame(f, fg_color="transparent")
            top.pack(fill="x", padx=14, pady=(12, 0))
            dot = ctk.CTkLabel(top, text="●", font=font(14), text_color=GREEN)
            dot.pack(side="left")
            title = ctk.CTkLabel(top, text="", font=font(15, "bold"), text_color=TEXT, anchor="w")
            title.pack(side="left", padx=6, fill="x", expand=True)
            sched = ctk.CTkLabel(f, text="", font=font(12), text_color=MUTED, anchor="w")
            sched.pack(fill="x", padx=14)
            nxt = ctk.CTkLabel(f, text="", font=font(12, "bold"), text_color=(GREEN, "#4ade80"), anchor="w")
            nxt.pack(fill="x", padx=14, pady=(0, 12))
            for w in (f, top, dot, title, sched, nxt):
                w.bind("<Button-1>", lambda _e, sid=s["id"]: self.select(sid))
            self.cards[s["id"]] = {"frame": f, "dot": dot, "title": title, "sched": sched, "next": nxt}
            self.refresh_card(s)
        self.highlight()

    def refresh_card(self, s):
        c = self.cards.get(s["id"])
        if not c:
            return
        c["title"].configure(text=host_of(s["url"]) or s["url"] or "New website")
        c["sched"].configure(text=schedule_text(s))
        c["dot"].configure(text_color=GREEN if s.get("enabled") and s.get("url") else "#9ca3af")

    def highlight(self):
        for sid, c in self.cards.items():
            c["frame"].configure(border_color=GREEN if sid == self.sel_id else CARD_BORDER)

    # ------------------------------------------------------------------ settings page
    def build_settings_page(self):
        page = SmoothScrollFrame(self.root, fg_color="transparent")
        page.grid_columnconfigure(0, weight=1)
        self.header(page, "Settings", "Chrome connection, startup and active hours.")

        self.v_startup = tk.IntVar(value=1 if self.cfg.get("startup", True) else 0)
        self.v_hours_on = tk.IntVar(value=1 if self.cfg.get("hours_on") else 0)
        self.v_from = tk.StringVar(value=self.cfg.get("hours_from", "09:00"))
        self.v_to = tk.StringVar(value=self.cfg.get("hours_to", "23:00"))
        self.v_chrome = tk.StringVar(value=self.cfg.get("chrome_path", ""))
        self.v_theme = tk.StringVar(value=self.cfg.get("theme", "System"))

        c = self.card(page, "Your Chrome")
        c.grid(row=1, column=0, sticky="ew", pady=(0, 12))
        ctk.CTkLabel(c, text="AutoBrowse works inside your normal Chrome, with your own logins, "
                             "through a small helper you install once.",
                     font=font(13), text_color=MUTED).pack(anchor="w", padx=18)
        r = self.row(c, (12, 4))
        self.label(r, "Chrome profile")
        self.profiles = chrome_profiles()
        self.profile_labels = ["%s  (%s)" % (n, f) if n != f else f for f, n in self.profiles]
        self.v_profile = tk.StringVar(value=self.profile_labels[0])
        current = self.cfg.get("profile") or "Default"
        for (folder, _), lab in zip(self.profiles, self.profile_labels):
            if folder == current:
                self.v_profile.set(lab)
        ctk.CTkOptionMenu(r, values=self.profile_labels, variable=self.v_profile, width=300, height=34,
                          corner_radius=8, fg_color=GREEN, button_color=GREEN_HOVER,
                          button_hover_color="#166534", font=font(13)).pack(side="left", padx=12)
        hs = self.row(c, (10, 4))
        self.helper_dot = ctk.CTkLabel(hs, text="●", font=font(16), text_color="#9ca3af")
        self.helper_dot.pack(side="left")
        self.helper_status = ctk.CTkLabel(hs, text="Checking...", font=font(13, "bold"), text_color=TEXT)
        self.helper_status.pack(side="left", padx=6)
        r = self.row(c, (10, 16))
        ctk.CTkButton(r, text="Install Chrome helper", height=38, corner_radius=10, font=font(13, "bold"),
                      fg_color=GREEN, hover_color=GREEN_HOVER, command=self.install_helper).pack(side="left")
        ctk.CTkButton(r, text="Hide Chrome's notice bar", height=38, corner_radius=10, font=font(13, "bold"),
                      fg_color="transparent", border_width=2, border_color=GREEN, text_color=(GREEN, "#4ade80"),
                      hover_color=("#dcfce7", "#173d27"), command=self.quiet_shortcuts).pack(side="left", padx=10)

        c = self.card(page, "General")
        c.grid(row=2, column=0, sticky="ew", pady=(0, 12))
        ctk.CTkSwitch(c, text="Start with Windows (runs quietly in the system tray)", variable=self.v_startup,
                      font=font(13), progress_color=GREEN).pack(anchor="w", padx=18, pady=(4, 8))
        self.v_busy = tk.IntVar(value=1 if self.cfg.get("pause_busy", True) else 0)
        self.v_busy_min = tk.StringVar(value=nice(self.cfg.get("busy_minutes", 2)))
        self.v_battery = tk.IntVar(value=1 if self.cfg.get("pause_battery") else 0)
        r = self.row(c)
        ctk.CTkCheckBox(r, text="Pause when I'm busy: wait until the keyboard and mouse are idle for",
                        variable=self.v_busy, font=font(13), fg_color=GREEN, hover_color=GREEN_HOVER).pack(side="left")
        self.small_entry(r, self.v_busy_min, 50)
        self.label(r, "minutes")
        r = self.row(c, (4, 8))
        ctk.CTkCheckBox(r, text="Pause while the laptop is on battery", variable=self.v_battery, font=font(13),
                        fg_color=GREEN, hover_color=GREEN_HOVER).pack(side="left")
        r = self.row(c)
        ctk.CTkSwitch(r, text="Only run between", variable=self.v_hours_on, font=font(13),
                      progress_color=GREEN).pack(side="left")
        self.small_entry(r, self.v_from, 70)
        self.label(r, "and")
        self.small_entry(r, self.v_to, 70)
        self.label(r, "(24 hour time)", muted=True)
        r = self.row(c, (12, 4))
        self.label(r, "Appearance")
        ctk.CTkSegmentedButton(r, values=["System", "Light", "Dark"], variable=self.v_theme, font=font(13),
                               selected_color=GREEN, selected_hover_color=GREEN_HOVER,
                               command=lambda v: ctk.set_appearance_mode(v)).pack(side="left", padx=12)
        ctk.CTkLabel(c, text="Chrome location (leave empty to find it automatically)", font=font(13),
                     text_color=MUTED).pack(anchor="w", padx=18, pady=(12, 2))
        r = self.row(c, (2, 16))
        ctk.CTkEntry(r, textvariable=self.v_chrome, height=34, corner_radius=8).pack(side="left", fill="x", expand=True)
        ctk.CTkButton(r, text="Browse", width=90, height=34, corner_radius=8, fg_color="transparent",
                      border_width=1, border_color=GREEN, text_color=(GREEN, "#4ade80"),
                      hover_color=("#dcfce7", "#173d27"), command=self.pick_chrome).pack(side="left", padx=(8, 0))

        r = ctk.CTkFrame(page, fg_color="transparent")
        r.grid(row=3, column=0, sticky="w", pady=(4, 10))
        ctk.CTkButton(r, text="Save", width=120, height=40, corner_radius=10, font=font(14, "bold"),
                      fg_color=GREEN, hover_color=GREEN_HOVER, command=self.save).pack(side="left")
        self.saved_label2 = ctk.CTkLabel(r, text="", font=font(13, "bold"), text_color=GREEN)
        self.saved_label2.pack(side="left", padx=12)
        return page

    # ------------------------------------------------------------------ activity page
    def build_activity_page(self):
        page = ctk.CTkFrame(self.root, fg_color="transparent")
        page.grid_columnconfigure(0, weight=1)
        page.grid_rowconfigure(2, weight=1)
        self.header(page, "Activity", "How long AutoBrowse has been active, and everything it did.")
        tiles = ctk.CTkFrame(page, fg_color="transparent")
        tiles.grid(row=1, column=0, sticky="ew", pady=(0, 14))
        self.stat_tiles = {}
        for i, (key, label) in enumerate((("since", "Since last reset"), ("today", "Today"), ("week", "Last 7 days"))):
            tiles.grid_columnconfigure(i, weight=1, uniform="t")
            t = self.card(tiles)
            t.grid(row=0, column=i, sticky="ew", padx=(0 if i == 0 else 6, 0 if i == 2 else 6))
            ctk.CTkLabel(t, text=label.upper(), font=font(11, "bold"), text_color=MUTED).pack(anchor="w", padx=16, pady=(12, 0))
            big = ctk.CTkLabel(t, text="0m 00s", font=font(26, "bold"), text_color=TEXT)
            big.pack(anchor="w", padx=16)
            small = ctk.CTkLabel(t, text="0 visits", font=font(12), text_color=MUTED)
            small.pack(anchor="w", padx=16, pady=(0, 12))
            self.stat_tiles[key] = (big, small)
        self.reset_note = ctk.CTkLabel(tiles, text="", font=font(12), text_color=MUTED)
        self.reset_note.grid(row=1, column=0, columnspan=2, sticky="w", pady=(8, 0))
        ctk.CTkButton(tiles, text="Reset counter", width=130, height=32, corner_radius=8, fg_color="transparent",
                      border_width=1, border_color=GREEN, text_color=(GREEN, "#4ade80"),
                      hover_color=("#dcfce7", "#173d27"), command=self.reset_active).grid(row=1, column=2, sticky="e", pady=(8, 0))
        c = self.card(page)
        c.grid(row=2, column=0, sticky="nsew")
        self.log_box = ctk.CTkTextbox(c, font=ctk.CTkFont(family="Consolas", size=13), wrap="word",
                                      fg_color="transparent", state="disabled")
        self.log_box.pack(fill="both", expand=True, padx=10, pady=10)
        ctk.CTkButton(page, text="Open log file", width=140, height=36, corner_radius=10, fg_color="transparent",
                      border_width=1, border_color=GREEN, text_color=(GREEN, "#4ade80"),
                      hover_color=("#dcfce7", "#173d27"),
                      command=lambda: self.open_path(LOG_PATH)).grid(row=3, column=0, sticky="w", pady=(12, 0))
        exp = ctk.CTkFrame(page, fg_color="transparent")
        exp.grid(row=3, column=0, sticky="e", pady=(12, 0))
        self.export_note = ctk.CTkLabel(exp, text="", font=font(12), text_color=MUTED)
        self.export_note.pack(side="left", padx=10)
        ctk.CTkButton(exp, text="Export Excel report", width=170, height=36, corner_radius=10, font=font(13, "bold"),
                      fg_color=GREEN, hover_color=GREEN_HOVER, command=self.export_excel).pack(side="left")
        self.shown_lines = -1
        return page

    # ------------------------------------------------------------------ form logic
    def site_by_id(self, sid):
        return next((s for s in self.cfg["sites"] if s["id"] == sid), None)

    def load_form(self, s):
        self.v_url.set(s["url"])
        self.v_enabled.set(1 if s["enabled"] else 0)
        self.v_mode.set("Fixed" if s["mode"] == "fixed" else "Random")
        self.v_fixed.set(nice(s["fixed_min"]))
        self.v_rmin.set(nice(s["rand_min"]))
        self.v_rmax.set(nice(s["rand_max"]))
        self.v_browse.set(1 if s["browse"] else 0)
        self.v_pmin.set(str(s["pages_min"]))
        self.v_pmax.set(str(s["pages_max"]))
        self.v_smin.set(nice(s["stay_min"]))
        self.v_smax.set(nice(s["stay_max"]))
        self.v_scroll.set(1 if s["scroll"] else 0)
        self.v_home.set(1 if s["home"] else 0)
        self.v_front.set(1 if s["front"] else 0)
        self.v_real.set(1 if s.get("real_mouse") else 0)
        self.v_al_change.set(1 if s.get("alert_change", True) else 0)
        self.v_al_logout.set(1 if s.get("alert_logout", True) else 0)
        self.v_al_down.set(1 if s.get("alert_down", True) else 0)
        self.v_sh_on.set(1 if s.get("hours_on") else 0)
        self.v_sh_from.set(s.get("hours_from", "09:00"))
        self.v_sh_to.set(s.get("hours_to", "23:00"))
        self.v_lim_on.set(1 if s.get("limit_on") else 0)
        self.v_lim_visits.set(str(s.get("limit_visits", 50)))
        self.v_lim_min.set(nice(s.get("limit_minutes", 60)))
        self.skip_box.delete("1.0", "end")
        self.skip_box.insert("1.0", s["skip"])
        self.targets_box.delete("1.0", "end")
        self.targets_box.insert("1.0", s.get("targets", ""))
        self.timing_rows()

    def commit_form(self):
        s = self.site_by_id(self.sel_id)
        if not s:
            return
        old_timing = (s["mode"], s["fixed_min"], s["rand_min"], s["rand_max"])
        s["url"] = normalize_url(self.v_url.get())
        s["enabled"] = bool(self.v_enabled.get())
        s["mode"] = "fixed" if self.v_mode.get() == "Fixed" else "random"
        s["fixed_min"] = to_num(self.v_fixed.get(), s["fixed_min"], 0.2)
        a = to_num(self.v_rmin.get(), s["rand_min"], 0.2)
        b = to_num(self.v_rmax.get(), s["rand_max"], 0.2)
        s["rand_min"], s["rand_max"] = min(a, b), max(a, b)
        s["browse"] = bool(self.v_browse.get())
        a = to_num(self.v_pmin.get(), s["pages_min"], 1, True)
        b = to_num(self.v_pmax.get(), s["pages_max"], 1, True)
        s["pages_min"], s["pages_max"] = min(a, b), max(a, b)
        a = to_num(self.v_smin.get(), s["stay_min"], 3)
        b = to_num(self.v_smax.get(), s["stay_max"], 3)
        s["stay_min"], s["stay_max"] = min(a, b), max(a, b)
        s["scroll"] = bool(self.v_scroll.get())
        s["home"] = bool(self.v_home.get())
        s["front"] = bool(self.v_front.get())
        s["real_mouse"] = bool(self.v_real.get())
        s["alert_change"] = bool(self.v_al_change.get())
        s["alert_logout"] = bool(self.v_al_logout.get())
        s["alert_down"] = bool(self.v_al_down.get())
        s["hours_on"] = bool(self.v_sh_on.get())
        s["hours_from"] = clock(self.v_sh_from.get(), s.get("hours_from", "09:00"))
        s["hours_to"] = clock(self.v_sh_to.get(), s.get("hours_to", "23:00"))
        s["limit_on"] = bool(self.v_lim_on.get())
        s["limit_visits"] = to_num(self.v_lim_visits.get(), s.get("limit_visits", 50), 0, True)
        s["limit_minutes"] = to_num(self.v_lim_min.get(), s.get("limit_minutes", 60), 0)
        s["skip"] = " ".join(self.skip_box.get("1.0", "end").split())
        s["targets"] = " ".join(self.targets_box.get("1.0", "end").split())
        if old_timing != (s["mode"], s["fixed_min"], s["rand_min"], s["rand_max"]):
            self.scheduler.replan(s)
        self.refresh_card(s)

    def select(self, sid):
        if sid != self.sel_id:
            self.commit_form()
        self.sel_id = sid
        s = self.site_by_id(sid)
        if s:
            self.load_form(s)
            self.show_editor(True)
        else:
            self.show_editor(False)
        self.highlight()

    def add_site(self):
        self.commit_form()
        s = new_site()
        self.cfg["sites"].append(s)
        self.render_cards()
        self.select(s["id"])
        self.show_page("sites")
        self.root.after(80, self.url_entry.focus_set)

    def remove_site(self):
        s = self.site_by_id(self.sel_id)
        if not s:
            return
        if not messagebox.askyesno(APP_NAME, "Remove %s?" % (host_of(s["url"]) or "this website")):
            return
        self.cfg["sites"].remove(s)
        self.sel_id = None
        save_config(self.cfg)
        self.render_cards()
        if self.cfg["sites"]:
            self.select(self.cfg["sites"][0]["id"])
        else:
            self.show_editor(False)

    def save(self, quiet=False):
        self.commit_form()
        self.cfg["startup"] = bool(self.v_startup.get())
        self.cfg["hours_on"] = bool(self.v_hours_on.get())
        self.cfg["hours_from"] = self.v_from.get().strip() or "09:00"
        self.cfg["hours_to"] = self.v_to.get().strip() or "23:00"
        self.cfg["chrome_path"] = self.v_chrome.get().strip()
        self.cfg["theme"] = self.v_theme.get()
        self.cfg["pause_busy"] = bool(self.v_busy.get())
        self.cfg["busy_minutes"] = to_num(self.v_busy_min.get(), 2, 0.5)
        self.cfg["pause_battery"] = bool(self.v_battery.get())
        for (folder, _), lab in zip(self.profiles, self.profile_labels):
            if lab == self.v_profile.get():
                self.cfg["profile"] = folder
        save_config(self.cfg)
        cur = self.site_by_id(self.sel_id)
        if cur and self.v_url.get() != cur["url"]:
            self.v_url.set(cur["url"])
        try:
            set_startup(self.cfg["startup"])
        except OSError as e:
            log.info("Could not update Windows startup: %s", e)
        if not quiet:
            for lab in (self.saved_label, self.saved_label2):
                lab.configure(text="Saved")
            self.root.after(2500, lambda: [l.configure(text="") for l in (self.saved_label, self.saved_label2)])

    def visit_now(self):
        s = self.site_by_id(self.sel_id)
        if not s:
            return
        self.save(quiet=True)
        if not s["url"]:
            messagebox.showinfo(APP_NAME, "Enter the website address first.")
            return
        self.scheduler.run_now(s["id"])
        log.info("%s   Visit requested", host_of(s["url"]))

    def pick_chrome(self):
        p = filedialog.askopenfilename(title="Find chrome.exe",
                                       filetypes=[("Chrome", "chrome.exe"), ("Programs", "*.exe")])
        if p:
            self.v_chrome.set(p)


    def install_helper(self):
        self.save(quiet=True)
        try:
            write_helper()
        except OSError as e:
            messagebox.showerror(APP_NAME, "Could not prepare the helper: %s" % e)
            return
        self.root.clipboard_clear()
        self.root.clipboard_append(HELPER_DIR)
        self.open_path(HELPER_DIR)
        try:
            Chrome(self.cfg).launch("chrome://extensions", wait=False)
        except Exception as e:
            log.info("Chrome   %s", e)
        messagebox.showinfo(APP_NAME,
                            "One time setup in Chrome:\n\n"
                            "1. Open chrome://extensions if it did not open by itself.\n"
                            "2. Turn on Developer mode (top right corner).\n"
                            "3. Click Load unpacked.\n"
                            "4. Paste the folder path (already copied) and click Select Folder.\n\n"
                            "The status in AutoBrowse turns green when the helper is connected.")

    def quiet_shortcuts(self):
        if os.name != "nt":
            return
        script = (
            "$w = New-Object -ComObject WScript.Shell; $n = 0;"
            "$dirs = @([Environment]::GetFolderPath('Desktop'), \"$env:APPDATA\\Microsoft\\Windows\\Start Menu\\Programs\","
            "\"$env:APPDATA\\Microsoft\\Internet Explorer\\Quick Launch\\User Pinned\\TaskBar\");"
            "foreach ($d in $dirs) { if (Test-Path $d) { Get-ChildItem $d -Filter *.lnk -Recurse -ErrorAction SilentlyContinue |"
            " ForEach-Object { try { $s = $w.CreateShortcut($_.FullName);"
            " if ($s.TargetPath -like '*\\chrome.exe' -and $s.Arguments -notlike '*silent-debugger-extension-api*') {"
            " $s.Arguments = ($s.Arguments + ' --silent-debugger-extension-api').Trim(); $s.Save(); $n++ } } catch {} } } };"
            "Write-Output $n"
        )
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
                           capture_output=True, text=True, creationflags=0x08000000)
        count = (r.stdout or "0").strip().splitlines()[-1] if (r.stdout or "").strip() else "0"
        messagebox.showinfo(APP_NAME,
                            "Updated %s Chrome shortcut(s).\n\n"
                            "Fully close Chrome once and reopen it from your shortcut. After that, Chrome no longer\n"
                            "shows the \"started debugging this browser\" bar while AutoBrowse is working.\n"
                            "When AutoBrowse opens Chrome itself, the bar is already hidden." % count)

    def open_path(self, path):
        try:
            os.startfile(path)
        except (AttributeError, OSError):
            pass

    # window and tray --------------------------------------------------------
    def show(self):
        self.root.deiconify()
        self.root.lift()
        self.root.attributes("-topmost", True)
        self.root.after(300, lambda: self.root.attributes("-topmost", False))
        self.root.focus_force()

    def hide(self):
        self.save(quiet=True)
        self.root.withdraw()
        if not self.hint_shown:
            self.hint_shown = True
            try:
                self.tray.notify("AutoBrowse keeps running in the system tray.", APP_NAME)
            except Exception:
                pass

    def toggle_pause(self):
        self.scheduler.paused = not self.scheduler.paused
        paused = self.scheduler.paused
        self.pause_btn.configure(text="Resume" if paused else "Pause")
        self.tray.icon = make_icon(64, paused)
        self.tray.title = APP_NAME + (" (paused)" if paused else "")
        try:
            self.tray.update_menu()
        except Exception:
            pass
        log.info("Paused" if paused else "Resumed")

    def quit(self):
        try:
            self.save(quiet=True)
        except Exception:
            pass
        self.scheduler.stop_event.set()
        try:
            self.tray.stop()
        except Exception:
            pass
        self.root.destroy()

    def listen(self):
        while True:
            try:
                conn, _ = self.server.accept()
                conn.close()
                self.events.put("show")
            except OSError:
                return

    def poll(self):
        try:
            while True:
                ev = self.events.get_nowait()
                if isinstance(ev, tuple) and ev[0] == "toast":
                    self.show_toast(ev[1], ev[2])
                    continue
                if ev == "show":
                    self.show()
                elif ev == "runall":
                    self.scheduler.run_now()
                    log.info("Visit all requested")
                elif ev == "toggle":
                    self.toggle_pause()
                elif ev == "exit":
                    self.quit()
                    return
        except queue.Empty:
            pass
        self.root.after(200, self.poll)

    def alert(self, title, body):
        """Thread-safe pop-up for alerts."""
        log.info("Alert   %s", title)
        self.events.put(("toast", title, body))

    def export_excel(self):
        try:
            path = export_report()
        except Exception as e:
            self.export_note.configure(text="Could not export: %s" % e, text_color=RED)
            return
        self.export_note.configure(text="Saved to Documents, AutoBrowse Reports", text_color=MUTED)
        log.info("Excel report saved: %s", path)
        self.open_path(path)

    def reset_active(self):
        ACTIVE.reset()
        self.update_active()

    def update_active(self):
        now = time.time()
        today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0).timestamp()
        for key, since in (("since", ACTIVE.reset_at), ("today", today), ("week", now - 7 * 86400)):
            secs, n = ACTIVE.total(since)
            big, small = self.stat_tiles[key]
            big.configure(text=span(secs))
            small.configure(text="%d visit%s" % (n, "" if n == 1 else "s"))
        d = datetime.fromtimestamp(ACTIVE.reset_at)
        self.reset_note.configure(text="Counting since %d %s, %s" % (d.day, d.strftime("%b"),
                                                                     d.strftime("%I:%M %p").lstrip("0")))

    def show_toast(self, title, body):
        if self.toast is not None:
            try:
                self.toast.destroy()
            except tk.TclError:
                pass
        t = tk.Toplevel(self.root)
        self.toast = t
        t.overrideredirect(True)
        t.attributes("-topmost", True)
        t.configure(bg=GREEN)
        inner = tk.Frame(t, bg="#ffffff", padx=16, pady=12)
        inner.pack(fill="both", expand=True, padx=(4, 0))
        tk.Label(inner, text=title, font=("Segoe UI Semibold", 11), fg="#12261a", bg="#ffffff",
                 anchor="w").pack(fill="x")
        tk.Label(inner, text=body, font=("Segoe UI", 10), fg="#12261a", bg="#ffffff", wraplength=360,
                 justify="left", anchor="w").pack(fill="x", pady=(6, 0))
        row = tk.Frame(inner, bg="#ffffff")
        row.pack(fill="x", pady=(10, 0))

        def close():
            try:
                t.destroy()
            except tk.TclError:
                pass
        tk.Button(row, text="Open AutoBrowse", relief="flat", bg="#e2f3e8", fg="#12261a", padx=10,
                  cursor="hand2", command=lambda: (self.show(), self.show_page("activity"), close())).pack(side="left")
        tk.Button(row, text="Close", relief="flat", bg="#ffffff", fg="#5d6f64", padx=10, cursor="hand2",
                  command=close).pack(side="right")
        t.update_idletasks()
        w, h = 400, t.winfo_reqheight()
        t.geometry("%dx%d+%d+%d" % (w, h, t.winfo_screenwidth() - w - 24, t.winfo_screenheight() - h - 70))
        t.after(12000, close)

    def tick(self):
        now = time.time()
        allowed = within_hours(self.cfg)
        for s in self.cfg["sites"]:
            c = self.cards.get(s["id"])
            if not c:
                continue
            if self.scheduler.running_id == s["id"]:
                nxt = "Visiting now"
            elif not s.get("enabled"):
                nxt = "Turned off"
            elif not s.get("url"):
                nxt = "Add the website address"
            elif self.scheduler.paused:
                nxt = "Paused"
            elif not allowed:
                nxt = "Outside active hours"
            elif self.scheduler.blocked.get(s["id"]):
                nxt = self.scheduler.blocked[s["id"]]
            elif self.scheduler.wait_reason:
                nxt = self.scheduler.wait_reason
            else:
                t = self.scheduler.next_run.get(s["id"])
                if t:
                    left = max(0, int(t - now))
                    nxt = "Next visit %s  (in %dm %02ds)" % (
                        datetime.fromtimestamp(t).strftime("%I:%M %p").lstrip("0"), left // 60, left % 60)
                else:
                    nxt = "Starting soon"
            if c["next"].cget("text") != nxt:
                c["next"].configure(text=nxt)

        connected = BRIDGE.connected()
        outdated = connected and BRIDGE.version != HELPER_MANIFEST["version"]
        self.helper_dot.configure(text_color="#f59e0b" if outdated else ("#22c55e" if connected else RED))
        if outdated:
            text = "Helper needs a quick update: open chrome://extensions and click the reload arrow on AutoBrowse Helper."
        elif connected:
            text = "Chrome helper connected"
        else:
            text = "Chrome helper not connected. Open Chrome, or install the helper below."
        self.helper_status.configure(text=text)

        lines = list(LOG_LINES)
        if len(lines) != self.shown_lines:
            self.log_box.configure(state="normal")
            self.log_box.delete("1.0", "end")
            self.log_box.insert("end", "\n".join(reversed(lines)) if lines else "Nothing yet.")
            self.log_box.configure(state="disabled")
            self.shown_lines = len(lines)

        if self.scheduler.paused:
            state, color = "Paused", "#fbbf24"
        elif not allowed:
            state, color = "Waiting", "#fbbf24"
        elif self.scheduler.wait_reason:
            state, color = "Waiting", "#fbbf24"
        else:
            state, color = "Running", "#4ade80"
        count = sum(1 for s in self.cfg["sites"] if s.get("enabled") and s.get("url"))
        self.status_text.configure(text=state)
        self.status_dot.configure(text_color=color)
        self.status_sub.configure(text="%d active website%s" % (count, "" if count == 1 else "s"))
        self.update_active()
        self.root.after(1000, self.tick)

    def start(self, show_window):
        if show_window:
            self.show()
        else:
            self.root.withdraw()
        self.root.mainloop()


# ----------------------------------------------------------------------------
# Entry point
# ----------------------------------------------------------------------------

def claim_instance():
    s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        s.bind(("127.0.0.1", INSTANCE_PORT))
        s.listen(2)
        return s
    except OSError:
        s.close()
        try:
            socket.create_connection(("127.0.0.1", INSTANCE_PORT), 2).close()
        except OSError:
            pass
        return None


def main():
    if os.name == "nt":
        try:
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("AutoBrowse.App")
        except Exception:
            pass
    server = claim_instance()
    if not server:
        return
    app = App(server)
    tray_only = "--tray" in sys.argv
    app.start(show_window=not tray_only or not app.cfg["sites"])


if __name__ == "__main__":
    main()
